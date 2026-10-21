#!/usr/bin/env python3
"""Build the financial-plan Excel model from plan-data.json.

Usage: build_model.py <plan-data.json> <output.xlsx>

Every calculation is an Excel formula: the workbook is the single calculation
engine. Run recalc.py afterwards to compute the values (needed before
check_model.py or the Word generator can read them).
"""
import calendar
import json
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import plugin_env  # noqa: E402

plugin_env.ensure()

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import LineChart, Reference  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter as CL  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from i18n import LABELS, SHEETS  # noqa: E402

FONT = "Arial"
F_IN = Font(name=FONT, color="0000FF")
F_CALC = Font(name=FONT, color="000000")
F_LINK = Font(name=FONT, color="008000")
F_BOLD = Font(name=FONT, bold=True)
F_TITLE = Font(name=FONT, bold=True, size=14)
F_HDR = Font(name=FONT, bold=True, color="FFFFFF")
FILL_HDR = PatternFill("solid", start_color="1F4E5A")
FILL_SEC = PatternFill("solid", start_color="D9E2E5")
FILL_DEFAULT = PatternFill("solid", start_color="FFFF00")
FILL_TOTAL = PatternFill("solid", start_color="F2F2F2")
TOP = Border(top=Side(style="thin"))

EUR = '€ #,##0;(€ #,##0);"-"'
EUR2 = '€ #,##0.00;(€ #,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0.##;(#,##0.##);"-"'
DATE = "mmm-yy"
DATE_FULL = "dd/mm/yyyy"

PARAM_KEYS = {
    "p_VATPeriod": "company.vat_period_months", "p_TaxStd": "rules.corporate_tax_standard",
    "p_TaxRed": "rules.corporate_tax_reduced", "p_TaxRedCap": "rules.corporate_tax_reduced_bracket",
    "p_TaxRedMinRem": "rules.reduced_rate_min_director_pay", "p_TaxRedFrom": "rules.reduced_rate_pay_condition_from_period",
    "p_UseRed": "tax.reduced_rate", "p_Unavail": "financing.unavailable_equity",
    "p_Prepay": "tax.prepayments", "p_TaxLag": "tax.payment_lag_months", "p_SSC": "rules.employer_ssc_rate",
    "p_SalMonths": "rules.salary_months_per_year", "p_WageIdx": "rules.wage_indexation",
    "p_DepLag": "rules.depreciation_start_lag_months", "p_CashBuffer": "rules.cash_buffer_months",
}
RULES = json.loads((Path(__file__).parent.parent / "references/belgian-rules.json").read_text())
SPARE = {"costs": 5, "staff": 3, "investments": 3, "equity": 2, "loans": 2, "grants": 2, "dividends": 2}


def q(sheet):
    """Sheet prefix for formulas; apostrophes in names must be doubled."""
    return "'" + sheet.replace("'", "''") + "'!"


def ym(s):
    y, m = s[:7].split("-")
    return int(y), int(m)


def add_months(y, m, k):
    t = y * 12 + (m - 1) + k
    return t // 12, t % 12 + 1


def months_between(a, b):
    return (b[0] - a[0]) * 12 + b[1] - a[1]


def to_date(s):
    if not s:
        return None
    y, m = ym(s)
    d = int(s[8:10]) if len(s) >= 10 else 1
    return date(y, m, d)


class Model:
    def __init__(self, data):
        self.d = data
        self.lang = data["meta"]["language"]
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        c = data["company"]
        self.start = ym(c["start_date"])
        fy1 = ym(c["first_fy_end"])
        self.nfy = int(c.get("fiscal_years", 3))
        self.fy_end = [add_months(*fy1, 12 * i) for i in range(self.nfy)]
        self.n = months_between(self.start, self.fy_end[-1]) + 1
        self.ext = self.n + 12
        self.mc0 = 4 + self.nfy + 1
        self.months = [add_months(*self.start, k) for k in range(self.ext)]
        self.fy_of = []
        for (y, m) in self.months:
            fy = next((i + 1 for i, e in enumerate(self.fy_end) if months_between((y, m), e) >= 0), 0)
            self.fy_of.append(fy)
        self.rows = {}
        self.lists = {}  # row lists for the document generator, stored in _map meta

    # ---- helpers -------------------------------------------------------
    def T(self, key, **kw):
        s = LABELS[key][self.lang]
        return s.format(**kw) if kw else s

    def S(self, key):
        return SHEETS[key][self.lang]

    def R(self, key):
        return q(self.S(key))

    def col(self, k):
        return CL(self.mc0 + k)

    def span(self, r, n=None, sheet=None, absolute=True):
        n = n or self.n
        p = self.R(sheet) if sheet else ""
        a = "$" if absolute else ""
        return f"{p}{a}{self.col(0)}{a}{r}:{a}{self.col(n - 1)}{a}{r}"

    def mdiff(self, c, ref):
        return f"((YEAR({c}$2)-YEAR({ref}))*12+MONTH({c}$2)-MONTH({ref}))"

    def fy_label(self, i):
        start = self.start if i == 0 else add_months(*self.fy_end[i - 1], 1)
        e = self.fy_end[i]
        s_txt = self.d["company"]["start_date"] if i == 0 else f"{start[0]}-{start[1]:02d}-01"
        sd = to_date(s_txt)
        ed = date(e[0], e[1], calendar.monthrange(e[0], e[1])[1])
        return f"{self.T('fy_short')} {i + 1}", f"{sd:%d/%m/%Y} – {ed:%d/%m/%Y}"

    def name(self, nm, sheet, cell):
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=f"{self.R(sheet)}{cell}")

    def monthly_sheet(self, key, title, ncols=None):
        ncols = ncols or self.n
        ws = self.wb.create_sheet(self.S(key))
        ws["A1"] = title
        ws["A1"].font = F_TITLE
        ws["A2"] = self.T("month")
        ws["A3"] = self.T("fy_short")
        for k in range(ncols):
            c = self.col(k)
            ws[f"{c}2"] = f"={self.R('timeline')}{c}3"
            ws[f"{c}2"].number_format = DATE
            ws[f"{c}3"] = f"={self.R('timeline')}{c}4"
            for r in (2, 3):
                ws[f"{c}{r}"].font = F_HDR
                ws[f"{c}{r}"].fill = FILL_HDR
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
            ws.column_dimensions[c].width = 11
        for i in range(self.nfy):
            c = CL(4 + i)
            ws[f"{c}2"] = self.fy_label(i)[0]
            ws[f"{c}2"].font = F_HDR
            ws[f"{c}2"].fill = FILL_HDR
            ws[f"{c}3"] = i + 1
            ws[f"{c}3"].font = Font(name=FONT, color="FFFFFF")
            ws[f"{c}3"].fill = FILL_HDR
            ws.column_dimensions[c].width = 13
        for r in (2, 3):
            for c in "ABC":
                ws[f"{c}{r}"].fill = FILL_HDR
                ws[f"{c}{r}"].font = F_HDR
        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions[CL(4 + self.nfy)].width = 2
        ws.freeze_panes = f"{self.col(0)}4"
        ws._next = 5
        return ws

    def put(self, ws, label, f, kind="flow", param=None, pfmt=None, note=None, fmt=EUR, bold=False,
            ncols=None, key=None, font=None, pfont=F_IN, default=False, pcomment=None):
        """Write one monthly row. f(k, c) returns a formula/value for month k in column c."""
        ncols = ncols or self.n
        r = ws._next
        ws._next += 1
        ws.cell(r, 1, label).font = F_BOLD if bold else F_CALC
        if param is not None:
            pc = ws.cell(r, 2, param)
            pc.font = pfont
            pc.number_format = pfmt or NUM
            if default:
                pc.fill = FILL_DEFAULT
            if pcomment:
                pc.comment = Comment(pcomment, "fp-model")
        if note is not None:
            ws.cell(r, 3, note).font = Font(name=FONT, italic=True, color="666666")
        for k in range(ncols):
            c = self.col(k)
            v = f(k, c)
            cell = ws[f"{c}{r}"]
            cell.value = v
            cell.number_format = fmt
            if font:
                cell.font = font
            elif isinstance(v, str) and v.startswith("="):
                cell.font = Font(name=FONT, bold=bold)
            else:
                cell.font = F_IN
            if bold:
                cell.fill = FILL_TOTAL
        for i in range(self.nfy):
            c = CL(4 + i)
            if kind == "flow":
                ws[f"{c}{r}"] = f"=SUMIFS({self.span(r)},{self.span(3)},{i + 1})"
            elif kind == "stock":
                ws[f"{c}{r}"] = f"=SUMIFS({self.span(r)},{self.span(5, sheet='timeline')},{i + 1})"
            else:
                continue
            ws[f"{c}{r}"].number_format = fmt
            ws[f"{c}{r}"].font = Font(name=FONT, bold=True)
            ws[f"{c}{r}"].fill = FILL_TOTAL
        if key:
            self.rows[key] = (ws.title, r)
        return r

    def section(self, ws, text):
        r = ws._next
        ws._next += 1
        ws.cell(r, 1, text).font = F_BOLD
        for cc in range(1, self.mc0 + self.n):
            ws.cell(r, cc).fill = FILL_SEC
        return r

    def skip(self, ws, k=1):
        ws._next += k

    def ref(self, key, c):
        sheet, r = self.rows[key]
        return f"{q(sheet)}{c}{r}"

    def ref_row(self, key, ncols=None):
        sheet, r = self.rows[key]
        n = ncols or self.n
        return f"{q(sheet)}${self.col(0)}${r}:${self.col(n - 1)}${r}"

    # ---- sheets ----------------------------------------------------------
    def build(self):
        self.readme()
        self.params()
        self.inputs()
        self.timeline()
        self.revenue()
        self.costs()
        self.investments()
        self.financing()
        self.calc()
        self.cashflow()
        self.pl()
        self.balance()
        self.kpis()
        self.checks()
        order = ["readme", "kpi", "checks", "params", "inputs", "revenue", "costs", "invest", "finance",
                 "calc", "cf", "pl", "bs", "timeline"]
        self.wb._sheets = [self.wb[self.S(k)] for k in order]
        self.map_sheet()
        self.wb.calculation.fullCalcOnLoad = True
        return self.wb

    def map_sheet(self):
        """Hidden index of key rows, so scripts can find cells without depending on labels."""
        ws = self.wb.create_sheet("_map")
        ws.sheet_state = "hidden"
        meta = {"lang": self.lang, "nfy": self.nfy, "n": self.n, "first_month_col": self.mc0,
                "sheets": json.dumps({k: self.S(k) for k in SHEETS}),
                "fy_labels": json.dumps([self.fy_label(i) for i in range(self.nfy)]),
                "rules_checked": RULES["last_checked"], "rules_review_days": RULES["review_after_days"],
                "rules_verify": json.dumps(self.rules_to_verify()),
                "rules_used": json.dumps(sorted(self.rules_used())),
                "lists": json.dumps(self.lists, ensure_ascii=False)}
        r = 1
        for k, v in meta.items():
            ws.cell(r, 1, "meta"), ws.cell(r, 2, k), ws.cell(r, 3, v)
            r += 1
        for k, v in self.rows.items():
            ws.cell(r, 1, "row"), ws.cell(r, 2, k), ws.cell(r, 3, v[0]), ws.cell(r, 4, v[1])
            if len(v) > 2:
                ws.cell(r, 5, v[2])
            r += 1

    def readme(self):
        ws = self.wb.create_sheet(self.S("readme"))
        ws.column_dimensions["A"].width = 110
        c = self.d["company"]
        lines = [(self.T("readme_title", company=c["name"]), F_TITLE)]
        for k in ("readme_1", "readme_2", "readme_3", "readme_4", "readme_5", "readme_6", "readme_7"):
            lines.append((self.T(k), F_CALC))
        lines.append(("", F_CALC))
        lines.append((self.T("disclaimer_title"), F_BOLD))
        for k in ("disclaimer_1", "disclaimer_2", "disclaimer_3"):
            lines.append((self.T(k, date=self.fmt_date(RULES["last_checked"]), sheet=self.S("params")), F_CALC))
        lines.append(("", F_CALC))
        lines.append((self.T("readme_colors"), F_BOLD))
        lines.append((self.T("readme_blue"), F_IN))
        lines.append((self.T("readme_black"), F_CALC))
        lines.append((self.T("readme_green"), F_LINK))
        lines.append((self.T("readme_yellow"), F_CALC))
        lines.append(("", F_CALC))
        lines.append((self.T("readme_simpl"), F_BOLD))
        for k in ("simpl_1", "simpl_2", "simpl_3", "simpl_4", "simpl_5"):
            lines.append((self.T(k), F_CALC))
        for i, (t, f) in enumerate(lines, 1):
            ws.cell(i, 1, t).font = f
            ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")
            if t == self.T("readme_yellow"):
                ws.cell(i, 1).fill = FILL_DEFAULT

    def rule(self, key):
        return self.d.get("rules", {}).get(key, RULES["rules"][key]["value"])

    def price_path(self, s):
        """Monthly prices of a stream with a price_schedule: price_excl_vat until the first step, then the last step."""
        steps = sorted((ym(m), p) for m, p in s["price_schedule"].items())
        out, p = [], s["price_excl_vat"]
        for mo in self.months[:self.n]:
            for m, v in steps:
                if m == mo:
                    p = v
            out.append(p)
        return out

    def rules_used(self):
        """Registry rules that change a number in this plan; the document only lists these as 'to verify'."""
        staff = self.d.get("staff", [])
        used = {"corporate_tax_standard", "corporate_tax_reduced", "corporate_tax_reduced_bracket", "cash_buffer_months"}
        if (self.d.get("tax") or {}).get("reduced_rate") and self.nfy >= self.rule("reduced_rate_pay_condition_from_period"):
            used |= {"reduced_rate_min_director_pay", "reduced_rate_pay_condition_from_period"}
        if staff:
            used.add("wage_indexation")
        if any(not s.get("employer_cost_monthly") for s in staff):
            used |= {"employer_ssc_rate", "salary_months_per_year"}
        if self.d.get("investments"):
            used.add("depreciation_start_lag_months")
        return used

    def rules_to_verify(self):
        """Registry rules marked 'verify' that this plan uses unchanged."""
        own, used = self.d.get("rules", {}), self.rules_used()
        return [k for k, v in RULES["rules"].items() if v["status"] == "verify" and k not in own and k in used]

    def fmt_date(self, s):
        return to_date(s).strftime("%d/%m/%Y")

    def params(self):
        ws = self.wb.create_sheet(self.S("params"))
        ws["A1"] = self.T("params_title")
        ws["A1"].font = F_TITLE
        ws["A2"] = self.T("params_hint", date=self.fmt_date(RULES["last_checked"]))
        ws["A2"].font = Font(name=FONT, italic=True, color="666666")
        for col, w in zip("ABCDE", (52, 18, 70, 48, 26)):
            ws.column_dimensions[col].width = w
        c = self.d["company"]
        tax = self.d.get("tax", {})
        own = self.d.get("rules", {})
        own_src = self.d.get("rules_sources", {})
        dflt = set(self.d.get("defaults_used", []))
        # (name, label, value, format, note, rule key or source key)
        items = [
            ("p_Company", "company_name", c["name"], "@", None, None),
            ("p_Start", "start_date", to_date(c["start_date"]), DATE_FULL, None, None),
            ("p_FY1End", "first_fy_end", to_date(c["first_fy_end"]), DATE_FULL, None, None),
            ("p_NFY", "n_fy", self.nfy, "0", None, None),
            ("p_VATPeriod", "vat_period", c.get("vat_period_months", 3), "0", "src_vat_period", "vat"),
            ("p_TaxStd", "tax_std", self.rule("corporate_tax_standard"), PCT, "src_tax", "corporate_tax_standard"),
            ("p_TaxRed", "tax_red", self.rule("corporate_tax_reduced"), PCT, "src_tax", "corporate_tax_reduced"),
            ("p_TaxRedCap", "tax_red_cap", self.rule("corporate_tax_reduced_bracket"), EUR, "src_tax", "corporate_tax_reduced_bracket"),
            ("p_TaxRedMinRem", "tax_red_minrem", self.rule("reduced_rate_min_director_pay"), EUR, "src_tax", "reduced_rate_min_director_pay"),
            ("p_TaxRedFrom", "tax_red_from", self.rule("reduced_rate_pay_condition_from_period"), "0", "src_tax_red_from",
             "reduced_rate_pay_condition_from_period"),
            ("p_UseRed", "use_red", 1 if tax.get("reduced_rate") else 0, "0", "src_use_red", "wib"),
            ("p_Prepay", "prepay", 1 if tax.get("prepayments") else 0, "0", "src_prepay", "wib"),
            ("p_TaxLag", "tax_lag", tax.get("payment_lag_months", 9), "0", "src_tax_lag", "practice"),
            ("p_SSC", "ssc", self.rule("employer_ssc_rate"), PCT, "src_ssc", "employer_ssc_rate"),
            ("p_SalMonths", "sal_months", self.rule("salary_months_per_year"), "0.00", "src_sal_months", "salary_months_per_year"),
            ("p_WageIdx", "wage_idx", self.rule("wage_indexation"), PCT, "src_wage_idx", "wage_indexation"),
            ("p_DepLag", "dep_lag", self.rule("depreciation_start_lag_months"), "0", "src_dep_lag", "depreciation_start_lag_months"),
            ("p_CashBuffer", "cash_buffer", self.rule("cash_buffer_months"), "0.0", "src_cash_buffer", "cash_buffer_months"),
            ("p_Unavail", "unavail", (self.d.get("financing") or {}).get("unavailable_equity", 0), EUR, "src_unavail", None),
        ]
        r = 3
        for h, txt in enumerate((self.T("param"), self.T("value"), self.T("source_note"), self.T("official_source"), self.T("checked_on")), 1):
            ws.cell(r, h, txt).font = F_HDR
            ws.cell(r, h).fill = FILL_HDR
        grey = Font(name=FONT, italic=True, color="666666")
        for nm, lab, val, fmt, note, ref in items:
            r += 1
            ws.cell(r, 1, self.T(lab)).font = F_CALC
            cell = ws.cell(r, 2, val)
            cell.font = F_IN
            cell.number_format = fmt
            if PARAM_KEYS.get(nm) in dflt:
                cell.fill = FILL_DEFAULT
            if note:
                ws.cell(r, 3, self.T(note)).font = grey
            self.name(nm, "params", f"$B${r}")
            if not ref:
                continue
            rule = RULES["rules"].get(ref)
            if rule and ref in own:
                o = own_src.get(ref, {})
                ws.cell(r, 4, o.get("source") or self.T("own_value")).font = grey
                ws.cell(r, 5, self.fmt_date(o["checked"]) if o.get("checked") else "").font = grey
                continue
            src = RULES["sources"][rule["source"] if rule else ref]
            sc = ws.cell(r, 4, src["name"][self.lang])
            if src["url"]:
                sc.hyperlink = src["url"]
                sc.font = Font(name=FONT, color="0563C1", underline="single")
            else:
                sc.font = grey
            if rule:
                dc = ws.cell(r, 5, self.fmt_date(rule["checked"]))
                dc.font = grey
                if rule["status"] == "verify":
                    dc.value = f"{dc.value} – {self.T('to_verify')}"
                    dc.font = Font(name=FONT, bold=True, color="C55A11")

    def inputs(self):
        ws = self.wb.create_sheet(self.S("inputs"))
        ws["A1"] = self.T("inputs_title")
        ws["A1"].font = F_TITLE
        ws["A2"] = self.T("inputs_hint")
        ws["A2"].font = Font(name=FONT, italic=True, color="666666")
        widths = [40, 16, 14, 14, 14, 14, 12, 12, 14, 50]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[CL(i)].width = w
        self.inp = {}
        r = 4
        dflt = set(self.d.get("defaults_used", []))

        def table(key, headers, items, writer, spare, validations=()):
            nonlocal r
            ws.cell(r, 1, self.T(f"tbl_{key}")).font = F_BOLD
            for cc in range(1, len(widths) + 1):
                ws.cell(r, cc).fill = FILL_SEC
            r += 1
            for i, h in enumerate(headers, 1):
                ws.cell(r, i, self.T(h)).font = F_HDR
                ws.cell(r, i).fill = FILL_HDR
                ws.cell(r, i).alignment = Alignment(wrap_text=True)
            r += 1
            rows = []
            for it in list(items) + [None] * spare:
                vals = writer(it) if it else [(None, None)] * len(headers)
                for i, (v, fmt) in enumerate(vals, 1):
                    cell = ws.cell(r, i, v)
                    cell.font = F_IN
                    if fmt:
                        cell.number_format = fmt
                if it and it.get("id") in dflt:
                    for i in range(1, len(headers) + 1):
                        ws.cell(r, i).fill = FILL_DEFAULT
                rows.append((r, it))
                r += 1
            for col_letter, options in validations:
                dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{col_letter}{rows[0][0]}:{col_letter}{rows[-1][0]}")
            self.inp[key] = rows
            r += 1

        fin = self.d.get("financing", {})
        table("costs", ["h_name", "h_category", "h_amount", "h_interval", "h_start", "h_end", "h_vat", "h_index", "h_blank", "h_source"],
              self.d.get("cost_lines", []),
              lambda it: [(it["name"], None), (it["category"], "0"),
                          (None if it.get("schedule") else it["amount"], EUR2),
                          (self.T("manual") if it.get("schedule") else it.get("interval_months", 1), "0"),
                          (to_date(it.get("start")), DATE_FULL), (to_date(it.get("end")), DATE_FULL),
                          (it.get("vat_rate", 0.21), PCT), (it.get("indexation", 0), PCT), (None, None),
                          (it.get("source", ""), None)],
              SPARE["costs"], [("B", "61,618,640,65")])
        table("staff", ["h_role", "h_contract", "h_start", "h_end", "h_fte", "h_gross", "h_override", "h_extras", "h_blank", "h_source"],
              self.d.get("staff", []),
              lambda it: [(it["role"], None), (it.get("contract", ""), None), (to_date(it["start"]), DATE_FULL),
                          (to_date(it.get("end")), DATE_FULL), (it.get("fte", 1), "0.00"),
                          (it.get("gross_monthly"), EUR), (it.get("employer_cost_monthly"), EUR),
                          (it.get("extras_monthly", 0), EUR), (None, None), (it.get("source", ""), None)],
              SPARE["staff"])
        table("investments", ["h_name", "h_inv_cat", "h_amount", "h_date", "h_life", "h_vat", "h_inkind", "h_blank", "h_blank", "h_source"],
              self.d.get("investments", []),
              lambda it: [(it["name"], None), (it.get("category", 22), "0"), (it["amount"], EUR),
                          (to_date(it["date"]), DATE_FULL), (it.get("life_years", 3), "0.0"),
                          (it.get("vat_rate", 0.21), PCT), (1 if it.get("in_kind") else 0, "0"), (None, None),
                          (None, None), (it.get("source", ""), None)],
              SPARE["investments"], [("B", "21,22,28"), ("G", "0,1")])
        table("equity", ["h_name", "h_amount", "h_date", "h_blank", "h_blank", "h_blank", "h_blank", "h_blank", "h_blank", "h_source"],
              fin.get("equity", []),
              lambda it: [(it["name"], None), (it["amount"], EUR), (to_date(it["date"]), DATE_FULL)] + [(None, None)] * 6
              + [(it.get("source", ""), None)],
              SPARE["equity"])
        lt = f"{self.T('loan_bank')},{self.T('loan_founder')},{self.T('loan_other')}"
        table("loans", ["h_lender", "h_loan_type", "h_amount", "h_date", "h_rate", "h_term", "h_grace", "h_blank", "h_guarantee", "h_source"],
              fin.get("loans", []),
              lambda it: [(it["lender"], None), (self.T("loan_" + it.get("type", "bank")), None), (it["amount"], EUR),
                          (to_date(it["date"]), DATE_FULL), (it.get("rate", 0), PCT), (it["term_months"], "0"),
                          (it.get("grace_months", 0), "0"), (None, None), (it.get("guarantees", ""), None),
                          (it.get("source", ""), None)],
              SPARE["loans"], [("B", lt)])
        table("grants", ["h_name", "h_amount", "h_date", "h_blank", "h_blank", "h_blank", "h_blank", "h_blank", "h_blank", "h_source"],
              fin.get("grants", []),
              lambda it: [(it["name"], None), (it["amount"], EUR), (to_date(it["date"]), DATE_FULL)] + [(None, None)] * 6
              + [(it.get("source", ""), None)],
              SPARE["grants"])
        table("dividends", ["h_name", "h_div_fy", "h_amount", "h_date", "h_blank", "h_blank", "h_blank", "h_blank", "h_blank", "h_source"],
              fin.get("dividends", []),
              lambda it: [(it.get("name") or self.T("dividend"), None), (it["fy"], "0"), (it["amount"], EUR),
                          (to_date(it["date"]), DATE_FULL)] + [(None, None)] * 5 + [(it.get("source", ""), None)],
              SPARE["dividends"])

    def timeline(self):
        ws = self.monthly_sheet("timeline", self.T("timeline_title"), self.ext)
        for k in range(self.ext):
            c = self.col(k)
            y, m = self.months[k]
            ws[f"{c}2"] = k + 1
            ws[f"{c}2"].number_format = "0"
            ws[f"{c}3"] = date(y, m, 1)
            ws[f"{c}3"].number_format = DATE
            ws[f"{c}4"] = self.fy_of[k]
            fy_end = self.fy_of[k] if (k + 1 < self.ext and self.fy_of[k + 1] != self.fy_of[k] and self.fy_of[k]) else 0
            ws[f"{c}5"] = fy_end
            ws[f"{c}6"] = f"=IF(MOD(MONTH({c}3),p_VATPeriod)=0,1,0)"
            for r in (3, 4, 5):
                ws[f"{c}{r}"].font = Font(name=FONT, color="808080")
        for i in range(self.nfy):
            ws[f"{CL(4 + i)}3"] = None
        ws["A2"], ws["A3"], ws["A4"], ws["A5"], ws["A6"] = (self.T("month_nr"), self.T("month"), self.T("fy_short"),
                                                            self.T("fy_end_flag"), self.T("vat_period_end"))
        for r in (2, 3):
            ws[f"A{r}"].font = F_HDR

    def revenue(self):
        ws = self.monthly_sheet("revenue", self.T("revenue_title"))
        dflt = set(self.d.get("defaults_used", []))
        streams = self.d.get("revenue_streams", [])
        keys = {k: [] for k in ("rev", "rev_incl", "vat_out", "rec", "cogs", "ovc", "ovc_incl", "ovc_vat",
                                "stock", "purch", "purch_vat", "purch_incl", "pay")}
        last = self.n - 1
        for s in streams:
            sid = s["id"]
            d = lambda f: f"{sid}.{f}" in dflt or sid in dflt  # noqa: E731
            self.section(ws, s["name"])
            vol = s.get("volumes", {})
            r_vol = self.put(ws, self.T("volume"), lambda k, c: vol.get(f"{self.months[k][0]}-{self.months[k][1]:02d}", 0),
                             kind="flow", note=s.get("unit", ""), fmt=NUM)
            r_p = ws._next
            r_idx = r_p + 1
            if s.get("price_schedule"):
                # a price step in a given month: monthly input prices, carried forward from the last step
                prices = self.price_path(s)
                r_price = self.put(ws, self.T("price_excl"), lambda k, c: prices[k], kind=None, param=s["price_excl_vat"],
                                   pfmt=EUR2, fmt=EUR2, note=self.T("manual"), default=d("price_excl_vat"))
            else:
                r_price = self.put(ws, self.T("price_excl"), lambda k, c: f"=$B${r_p}*(1+$B${r_idx})^({c}$3-1)",
                                   kind=None, param=s["price_excl_vat"], pfmt=EUR2, fmt=EUR2, default=d("price_excl_vat"))
            self.put(ws, self.T("price_index"), lambda k, c: None, kind=None, param=s.get("price_indexation", 0), pfmt=PCT)
            r_rev = self.put(ws, self.T("revenue_excl"), lambda k, c: f"={c}{r_vol}*{c}{r_price}", bold=True)
            keys["rev"].append(r_rev)
            self.lists.setdefault("streams", []).append([sid, s["name"], r_vol, r_rev])
            r_vat = ws._next
            self.put(ws, self.T("vat_rate"), lambda k, c: f"={c}{r_rev}*$B${r_vat}", param=s.get("vat_rate", 0.21), pfmt=PCT,
                     default=d("vat_rate"))
            keys["vat_out"].append(r_vat)
            r_lag = ws._next
            self.put(ws, self.T("receipts"),
                     lambda k, c: f"=IF({k + 1}-ROUND($B${r_lag},0)>=1,INDEX({self.span(r_rev)},{k + 1}-ROUND($B${r_lag},0)),0)*(1+$B${r_vat})",
                     param=s.get("payment_lag_months", 0), pfmt="0", note=self.T("lag_months"), default=d("payment_lag_months"))
            keys["rec"].append(r_lag)
            r_gp = ws._next
            r_gu = ws._next + 1
            self.put(ws, self.T("goods_pct"), lambda k, c: None, kind=None, param=s.get("goods_cost_pct", 0), pfmt=PCT,
                     default=d("goods_cost_pct"))
            self.put(ws, self.T("goods_unit"), lambda k, c: None, kind=None, param=s.get("goods_cost_per_unit", 0), pfmt=EUR2,
                     default=d("goods_cost_per_unit"))
            r_cogs = self.put(ws, self.T("cogs"),
                              lambda k, c: f"={c}{r_rev}*$B${r_gp}+{c}{r_vol}*$B${r_gu}*(1+$B${r_idx})^({c}$3-1)")
            keys["cogs"].append(r_cogs)
            r_op = ws._next
            r_ov = ws._next + 1
            r_ovc = self.put(ws, self.T("other_var"), lambda k, c: f"={c}{r_rev}*$B${r_op}", param=s.get("other_var_cost_pct", 0),
                             pfmt=PCT, default=d("other_var_cost_pct"))
            keys["ovc"].append(r_ovc)
            self.put(ws, self.T("other_var_vat"), lambda k, c: f"={c}{r_ovc}*$B${r_ov}", param=s.get("other_var_vat_rate", 0.21), pfmt=PCT)
            keys["ovc_vat"].append(r_ov)
            r_cov = ws._next
            # stock covers next month's cost of goods; after the last month, the same month a year earlier stands in
            r_stock = self.put(ws, self.T("stock_end"),
                               lambda k, c: f"=$B${r_cov}*{self.col(k + 1 if k < last else k - 11)}{r_cogs}", kind="stock",
                               param=s.get("stock_cover_months", 0), pfmt="0.0", note=self.T("months_cover"),
                               default=d("stock_cover_months"))
            keys["stock"].append(r_stock)
            r_pur = self.put(ws, self.T("purchases"),
                             lambda k, c: f"={c}{r_cogs}+{c}{r_stock}" + (f"-{self.col(k - 1)}{r_stock}" if k else ""))
            keys["purch"].append(r_pur)
            r_pv = ws._next
            self.put(ws, self.T("purchases_vat"), lambda k, c: f"={c}{r_pur}*$B${r_pv}", param=s.get("purchase_vat_rate", s.get("vat_rate", 0.21)),
                     pfmt=PCT)
            keys["purch_vat"].append(r_pv)
            r_sl = ws._next
            self.put(ws, self.T("supplier_pay"),
                     lambda k, c: f"=IF({k + 1}-ROUND($B${r_sl},0)>=1,INDEX({self.span(r_pur)},{k + 1}-ROUND($B${r_sl},0)),0)*(1+$B${r_pv})",
                     param=s.get("supplier_lag_months", 1 if (s.get("goods_cost_pct") or s.get("goods_cost_per_unit")) else 0),
                     pfmt="0", note=self.T("lag_months"), default=d("supplier_lag_months"))
            keys["pay"].append(r_sl)
            self.skip(ws)

        self.section(ws, self.T("totals"))

        def total(key, label, rows, bold=False, kind="flow", tmpl="{c}{r}"):
            return self.put(ws, label, lambda k, c: "=" + ("+".join(tmpl.format(c=c, r=r) for r in rows) if rows else "0"),
                            bold=bold, kind=kind, key=key)

        total("rev", self.T("revenue_excl"), keys["rev"], bold=True)
        total("vat_out", self.T("vat_out"), keys["vat_out"])
        r_ri = total("rev_incl", self.T("revenue_incl"), keys["rev"] + keys["vat_out"])
        r_rc = total("rec", self.T("receipts"), keys["rec"])
        total("cogs", self.T("cogs"), keys["cogs"], bold=True)
        total("ovc", self.T("other_var"), keys["ovc"])
        total("ovc_vat", self.T("other_var_vat"), keys["ovc_vat"])
        total("stock", self.T("stock_end"), keys["stock"], kind="stock")
        r_pu = total("purch", self.T("purchases"), keys["purch"])
        r_pv = total("purch_vat", self.T("purchases_vat"), keys["purch_vat"])
        r_pp = total("pay", self.T("supplier_pay"), keys["pay"])
        self.put(ws, self.T("receivables"), lambda k, c: f"=SUM(${self.col(0)}{r_ri}:{c}{r_ri})-SUM(${self.col(0)}{r_rc}:{c}{r_rc})",
                 kind="stock", key="receivables")
        self.put(ws, self.T("payables"),
                 lambda k, c: f"=SUM(${self.col(0)}{r_pu}:{c}{r_pu})+SUM(${self.col(0)}{r_pv}:{c}{r_pv})-SUM(${self.col(0)}{r_pp}:{c}{r_pp})",
                 kind="stock", key="payables")

    def costs(self):
        ws = self.monthly_sheet("costs", self.T("costs_title"))
        ws["B4"], ws["C4"] = self.T("h_category"), self.T("h_vat")
        ws["B4"].font = ws["C4"].font = F_BOLD
        I = self.R("inputs")
        self.section(ws, self.T("tbl_costs"))
        first = ws._next
        for r_in, it in self.inp["costs"]:
            r = ws._next
            if it and it.get("schedule"):
                sch = it["schedule"]
                f = lambda k, c: sch.get(f"{self.months[k][0]}-{self.months[k][1]:02d}", 0)  # noqa: E731
            else:
                def f(k, c, r_in=r_in):
                    m = self.mdiff(c, f"{I}$E${r_in}")
                    e = self.mdiff(c, f"{I}$F${r_in}")
                    return (f"=IF(N({I}$C${r_in})=0,0,IF(AND({m}>=0,OR({I}$F${r_in}=\"\",{e}<=0),"
                            f"IF({I}$D${r_in}=0,{m}=0,MOD({m},{I}$D${r_in})=0)),"
                            f"{I}$C${r_in}*(1+{I}$H${r_in})^({c}$3-1),0))")
            self.put(ws, f"={I}A{r_in}&\"\"", f)
            ws.cell(r, 1).font = F_LINK
            ws.cell(r, 2, f"=IF({I}B{r_in}=\"\",\"\",{I}B{r_in})").font = F_LINK
            ws.cell(r, 3, f"=N({I}G{r_in})").font = F_LINK
            ws.cell(r, 3).number_format = PCT
        last = ws._next - 1
        B, C = f"$B${first}:$B${last}", f"$C${first}:$C${last}"

        def rng(c):
            return f"{c}{first}:{c}{last}"
        self.skip(ws)
        self.put(ws, self.T("cat61"), lambda k, c: f"=SUMIF({B},61,{rng(c)})+SUMIF({B},618,{rng(c)})", bold=True, key="c61")
        self.put(ws, self.T("cat618"), lambda k, c: f"=SUMIF({B},618,{rng(c)})", key="c618")
        self.put(ws, self.T("cat640"), lambda k, c: f"=SUMIF({B},640,{rng(c)})", bold=True, key="c640")
        self.put(ws, self.T("cat65"), lambda k, c: f"=SUMIF({B},65,{rng(c)})", bold=True, key="c65")
        self.put(ws, self.T("cat61_paid"), lambda k, c: f"=SUMPRODUCT((({B}=61)+({B}=618))*{rng(c)}*(1+{C}))", key="c61_paid")
        self.put(ws, self.T("cat640_paid"), lambda k, c: f"=SUMPRODUCT(({B}=640)*{rng(c)}*(1+{C}))", key="c640_paid")
        self.put(ws, self.T("cat65_paid"), lambda k, c: f"=SUMPRODUCT(({B}=65)*{rng(c)}*(1+{C}))", key="c65_paid")
        self.put(ws, self.T("vat_in_costs"), lambda k, c: f"=SUMPRODUCT({rng(c)}*{C})", key="c_vat")
        self.rows["cost_lines"] = (ws.title, first, last)

        self.skip(ws)
        self.section(ws, self.T("tbl_staff"))
        cost_rows, fte_rows = [], []
        for r_in, it in self.inp["staff"]:
            def f(k, c, r_in=r_in):
                m = self.mdiff(c, f"{I}$C${r_in}")
                e = self.mdiff(c, f"{I}$D${r_in}")
                return (f"=IF(N({I}$E${r_in})+N({I}$G${r_in})=0,0,IF(AND({m}>=0,OR({I}$D${r_in}=\"\",{e}<=0)),"
                        f"IF(N({I}$G${r_in})>0,{I}$G${r_in},{I}$F${r_in}*{I}$E${r_in}*p_SalMonths/12*(1+p_SSC))"
                        f"*(1+p_WageIdx)^({c}$3-1)+N({I}$H${r_in}),0))")
            r = self.put(ws, f"={I}A{r_in}&\"\"", f)
            ws.cell(r, 1).font = F_LINK
            cost_rows.append(r)
            self.lists.setdefault("staff", []).append(r)

            def g(k, c, r_in=r_in):
                m = self.mdiff(c, f"{I}$C${r_in}")
                e = self.mdiff(c, f"{I}$D${r_in}")
                return f"=IF(AND(N({I}$E${r_in})>0,{m}>=0,OR({I}$D${r_in}=\"\",{e}<=0)),{I}$E${r_in},0)"
            r2 = self.put(ws, f"=\"  FTE \"&{I}A{r_in}", g, kind="stock", fmt="0.0")
            ws.cell(r2, 1).font = F_LINK
            fte_rows.append(r2)
        self.put(ws, self.T("cat62"), lambda k, c: "=" + "+".join(f"{c}{r}" for r in cost_rows), bold=True, key="c62")
        self.put(ws, self.T("fte_total"), lambda k, c: "=" + "+".join(f"{c}{r}" for r in fte_rows), kind="stock", fmt="0.0", key="fte")

    def investments(self):
        ws = self.monthly_sheet("invest", self.T("invest_title"))
        I = self.R("inputs")
        capex, dep, nbv = [], [], []
        for r_in, it in self.inp["investments"]:
            self.section(ws, f"={I}A{r_in}&\"\"")
            m = lambda c, r_in=r_in: self.mdiff(c, f"{I}$D${r_in}")  # noqa: E731
            rc = self.put(ws, self.T("capex"), lambda k, c, r_in=r_in: f"=IF({m(c)}=0,N({I}$C${r_in}),0)")
            rd = self.put(ws, self.T("depreciation"),
                          lambda k, c, r_in=r_in: f"=IF(N({I}$E${r_in})*N({I}$C${r_in})=0,0,IF(AND({m(c)}>=p_DepLag,{m(c)}<p_DepLag+{I}$E${r_in}*12),{I}$C${r_in}/({I}$E${r_in}*12),0))")
            rn = self.put(ws, self.T("nbv"),
                          lambda k, c, r_in=r_in, rd=rd: f"=IF({m(c)}>=0,N({I}$C${r_in}),0)-SUM(${self.col(0)}{rd}:{c}{rd})",
                          kind="stock")
            capex.append((rc, r_in))
            dep.append(rd)
            nbv.append((rn, r_in))
        self.skip(ws)
        self.section(ws, self.T("totals"))
        self.put(ws, self.T("capex"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r, _ in capex), bold=True, key="capex")
        self.put(ws, self.T("capex_paid"), lambda k, c: "=0" + "".join(f"+{c}{r}*(1+N({I}$F${ri}))*(1-N({I}$G${ri}))" for r, ri in capex),
                 key="capex_paid")
        self.put(ws, self.T("capex_vat"), lambda k, c: "=0" + "".join(f"+{c}{r}*N({I}$F${ri})*(1-N({I}$G${ri}))" for r, ri in capex),
                 key="capex_vat")
        self.put(ws, self.T("inkind"), lambda k, c: "=0" + "".join(f"+{c}{r}*N({I}$G${ri})" for r, ri in capex), key="inkind")
        self.put(ws, self.T("depreciation"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r in dep), bold=True, key="dep")
        for cat, key in ((21, "nbv21"), (22, "nbv22"), (28, "nbv28")):
            self.put(ws, self.T(key), lambda k, c, cat=cat: "=0" + "".join(f"+{c}{r}*({I}$B${ri}={cat})" for r, ri in nbv),
                     kind="stock", key=key)

    def financing(self):
        ws = self.monthly_sheet("finance", self.T("finance_title"), self.ext)
        I = self.R("inputs")
        n = self.ext
        self.section(ws, self.T("tbl_equity"))
        eq = []
        for r_in, it in self.inp["equity"]:
            r = self.put(ws, f"={I}A{r_in}&\"\"", lambda k, c, r_in=r_in: f"=IF({self.mdiff(c, f'{I}$C${r_in}')}=0,N({I}$B${r_in}),0)", ncols=n)
            ws.cell(r, 1).font = F_LINK
            eq.append(r)
        self.put(ws, self.T("equity_cash"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r in eq), bold=True, ncols=n, key="eq_cash")
        self.skip(ws)
        self.section(ws, self.T("tbl_grants"))
        gr = []
        for r_in, it in self.inp["grants"]:
            r = self.put(ws, f"={I}A{r_in}&\"\"", lambda k, c, r_in=r_in: f"=IF({self.mdiff(c, f'{I}$C${r_in}')}=0,N({I}$B${r_in}),0)", ncols=n)
            ws.cell(r, 1).font = F_LINK
            gr.append(r)
        self.put(ws, self.T("grants"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r in gr), bold=True, ncols=n, key="grants")
        self.skip(ws)
        # dividends: decided at the end of the financial year they come from (appropriation of its result),
        # or at payment when paid during that year (interim dividend)
        self.section(ws, self.T("tbl_dividends"))
        dd, dp = [], []
        for r_in, it in self.inp["dividends"]:
            decl = f"MIN({I}$D${r_in},EOMONTH(p_FY1End,12*(N({I}$B${r_in})-1)))"
            r = self.put(ws, f"={I}A{r_in}&\" ({self.T('div_decided')})\"",
                         lambda k, c, r_in=r_in, decl=decl: f"=IF(N({I}$C${r_in})=0,0,IF({self.mdiff(c, decl)}=0,{I}$C${r_in},0))", ncols=n)
            ws.cell(r, 1).font = F_LINK
            dd.append(r)
            r = self.put(ws, f"={I}A{r_in}&\" ({self.T('div_paid_lbl')})\"",
                         lambda k, c, r_in=r_in: f"=IF({self.mdiff(c, f'{I}$D${r_in}')}=0,N({I}$C${r_in}),0)", ncols=n)
            ws.cell(r, 1).font = F_LINK
            dp.append(r)
        r_dd = self.put(ws, self.T("div_decided_total"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r in dd), bold=True, ncols=n, key="div_decl")
        r_dp = self.put(ws, self.T("div_paid_total"), lambda k, c: "=0" + "".join(f"+{c}{r}" for r in dp), bold=True, ncols=n, key="div_paid")
        self.put(ws, self.T("div_payable"), lambda k, c: f"=SUM(${self.col(0)}{r_dd}:{c}{r_dd})-SUM(${self.col(0)}{r_dp}:{c}{r_dp})",
                 kind="stock", ncols=n, key="div_payable")
        self.skip(ws)
        loans = []
        for r_in, it in self.inp["loans"]:
            self.section(ws, f"={I}A{r_in}&\" (\"&{I}B{r_in}&\")\"")
            m = lambda c, r_in=r_in: self.mdiff(c, f"{I}$D${r_in}")  # noqa: E731
            rd = ws._next
            ri, rp, rb = rd + 1, rd + 2, rd + 3
            prev = lambda k, rb=rb: f"{self.col(k - 1)}{rb}" if k else "0"  # noqa: E731
            self.put(ws, self.T("loan_disb"), lambda k, c, r_in=r_in: f"=IF({m(c)}=0,N({I}$C${r_in}),0)", ncols=n)
            self.put(ws, self.T("loan_interest"), lambda k, c, r_in=r_in: f"={prev(k)}*N({I}$E${r_in})/12", ncols=n)
            self.put(ws, self.T("loan_principal"),
                     lambda k, c, r_in=r_in: (f"=IF(N({I}$C${r_in})=0,0,IF(AND({m(c)}>N({I}$G${r_in}),{m(c)}<={I}$F${r_in}),"
                                              f"MIN({prev(k)},PMT(N({I}$E${r_in})/12,{I}$F${r_in}-N({I}$G${r_in}),-{I}$C${r_in})-{c}{ri}),0))"),
                     ncols=n)
            self.put(ws, self.T("loan_balance"), lambda k, c: f"={prev(k)}+{c}{rd}-{c}{rp}", kind="stock", ncols=n)
            rcur = self.put(ws, self.T("loan_current"),
                            lambda k, c: f"=MIN({c}{rb},SUMIFS({self.span(rp, n)},{self.span(2, n)},\">\"&{c}$2,{self.span(2, n)},\"<=\"&EDATE({c}$2,12)))",
                            kind="stock", ncols=n)
            loans.append((rd, ri, rp, rb, rcur, r_in))
            self.skip(ws)
        self.section(ws, self.T("totals"))
        fnd = f"\"{self.T('loan_founder')}\""

        def tot(key, label, idx, filt=None, kind="flow", bold=False):
            def f(k, c):
                parts = []
                for L in loans:
                    w = "" if filt is None else (f"*({I}$B${L[5]}={fnd})" if filt == "founder" else f"*({I}$B${L[5]}<>{fnd})")
                    parts.append(f"+{c}{L[idx]}{w}")
                return "=0" + "".join(parts)
            return self.put(ws, label, f, kind=kind, ncols=n, key=key, bold=bold)
        tot("loan_disb", self.T("loan_disb"), 0, bold=True)
        tot("loan_int", self.T("loan_interest"), 1, bold=True)
        tot("loan_princ", self.T("loan_principal"), 2, bold=True)
        tot("loan_bal_fin", self.T("loan_bal_fin"), 3, "fin", "stock")
        tot("loan_bal_fnd", self.T("loan_bal_fnd"), 3, "founder", "stock")
        tot("loan_cur_fin", self.T("loan_cur_fin"), 4, "fin", "stock")
        tot("loan_cur_fnd", self.T("loan_cur_fnd"), 4, "founder", "stock")

    def calc(self):
        ws = self.monthly_sheet("calc", self.T("calc_title"))
        ws["A4"] = self.T("fy_end_flag")
        for k in range(self.n):
            c = self.col(k)
            ws[f"{c}4"] = f"={self.R('timeline')}{c}5"
            ws[f"{c}4"].font = F_LINK
        L = lambda key: (lambda k, c: f"={self.ref(key, c)}")  # noqa: E731
        self.section(ws, self.T("pl_monthly"))
        r70 = self.put(ws, self.T("pl_70"), L("rev"), key="m70")
        r74 = self.put(ws, self.T("pl_74"), L("grants"), key="m74")
        r60 = self.put(ws, self.T("pl_60"), L("cogs"), key="m60")
        r61 = self.put(ws, self.T("pl_61"), lambda k, c: f"={self.ref('c61', c)}+{self.ref('ovc', c)}", key="m61")
        r62 = self.put(ws, self.T("pl_62"), L("c62"), key="m62")
        r630 = self.put(ws, self.T("pl_630"), L("dep"), key="m630")
        r640 = self.put(ws, self.T("pl_640"), L("c640"), key="m640")
        r9901 = self.put(ws, self.T("pl_9901"), lambda k, c: f"={c}{r70}+{c}{r74}-{c}{r60}-{c}{r61}-{c}{r62}-{c}{r630}-{c}{r640}",
                         bold=True, key="m9901")
        r65 = self.put(ws, self.T("pl_65"), lambda k, c: f"={self.ref('loan_int', c)}+{self.ref('c65', c)}", key="m65")
        r9903 = self.put(ws, self.T("pl_9903"), lambda k, c: f"={c}{r9901}-{c}{r65}", bold=True, key="m9903")
        self.rows["m_tax"] = (ws.title, ws._next)
        r_tax = ws._next
        # filled after the P&L sheet exists (needs its tax row)
        self.put(ws, self.T("pl_67"), lambda k, c: None)
        r9904 = self.put(ws, self.T("pl_9904"), lambda k, c: f"={c}{r9903}-{c}{r_tax}", bold=True, key="m9904")
        self.put(ws, self.T("cum_result"), lambda k, c: f"=SUM(${self.col(0)}{r9904}:{c}{r9904})", kind="stock", key="cum_result")
        self.skip(ws)
        self.section(ws, self.T("vat_section"))
        r_out = self.put(ws, self.T("vat_out"), L("vat_out"))
        r_in = self.put(ws, self.T("vat_in"), lambda k, c: "=" + "+".join(self.ref(x, c) for x in ("purch_vat", "ovc_vat", "c_vat", "capex_vat")))
        r_pay = ws._next
        r_bal = ws._next + 1
        T = self.R("timeline")
        self.put(ws, self.T("vat_paid"),
                 lambda k, c: f"=IF({T}{self.col(k - 1)}6=1,MAX(0,{self.col(k - 1)}{r_bal}),0)" if k else "=0", key="vat_paid")
        self.put(ws, self.T("vat_balance"),
                 lambda k, c: (f"={self.col(k - 1)}{r_bal}+" if k else "=") + f"{c}{r_out}-{c}{r_in}-{c}{r_pay}", kind="stock", key="vat_bal")
        self.skip(ws)
        self.section(ws, self.T("tax_section"))
        r_tp = ws._next
        self.put(ws, self.T("tax_paid"),
                 lambda k, c: f"=IF(p_Prepay=1,{c}{r_tax},IF({k + 1}-p_TaxLag>=1,INDEX({self.span(r_tax)},{k + 1}-p_TaxLag),0))", key="tax_paid")
        self.put(ws, self.T("tax_payable"), lambda k, c: f"=SUM(${self.col(0)}{r_tax}:{c}{r_tax})-SUM(${self.col(0)}{r_tp}:{c}{r_tp})",
                 kind="stock", key="tax_payable")
        self.skip(ws)
        self.section(ws, self.T("equity_section"))
        r_contr = self.put(ws, self.T("contributions"), lambda k, c: f"={self.ref('eq_cash', c)}+{self.ref('inkind', c)}", key="contr")
        r_cc = self.put(ws, self.T("cum_contributions"), lambda k, c: f"=SUM(${self.col(0)}{r_contr}:{c}{r_contr})", kind="stock", key="cum_contr")
        _, r_cr = self.rows["cum_result"]
        r_cd = self.put(ws, self.T("cum_dividends"), lambda k, c: f"=SUM({self.R('finance')}${self.col(0)}${self.rows['div_decl'][1]}:{c}${self.rows['div_decl'][1]})",
                        kind="stock", key="cum_div")
        r_em = self.put(ws, self.T("equity_month"), lambda k, c: f"={c}{r_cc}+{c}{r_cr}-{c}{r_cd}", kind="stock", key="equity_m")
        # art. 5:142 WVV: no distribution when the net assets are, or would become, lower than the unavailable equity
        self.put(ws, self.T("div_test"), lambda k, c: f"=IF({self.ref('div_decl', c)}>0,IF({c}{r_em}<p_Unavail,1,0),0)", key="div_test")
        self.put(ws, self.T("fixed_costs"), lambda k, c: f"={self.ref('c61', c)}+{self.ref('c62', c)}+{self.ref('c640', c)}", key="fixed")

    def cashflow(self):
        ws = self.monthly_sheet("cf", self.T("cf_title"))
        L = lambda key, sign="": (lambda k, c: f"={sign}{self.ref(key, c)}")  # noqa: E731
        self.section(ws, self.T("cf_operating"))
        a = self.put(ws, self.T("cf_receipts"), L("rec"))
        self.put(ws, self.T("cf_grants"), L("grants"))
        self.put(ws, self.T("cf_suppliers"), L("pay", "-"))
        self.put(ws, self.T("cf_services"), lambda k, c: f"=-{self.ref('c61_paid', c)}-{self.ref('ovc', c)}-{self.ref('ovc_vat', c)}")
        self.put(ws, self.T("cf_personnel"), L("c62", "-"))
        self.put(ws, self.T("cf_other"), L("c640_paid", "-"))
        self.put(ws, self.T("cf_vat"), L("vat_paid", "-"))
        b = self.put(ws, self.T("cf_tax"), L("tax_paid", "-"))
        r_op = self.put(ws, self.T("cf_op_net"), lambda k, c: f"=SUM({c}{a}:{c}{b})", bold=True, key="cf_op")
        self.section(ws, self.T("cf_investing"))
        r_inv = self.put(ws, self.T("cf_capex"), L("capex_paid", "-"), bold=True, key="cf_inv")
        self.section(ws, self.T("cf_financing"))
        a = self.put(ws, self.T("cf_equity"), L("eq_cash"))
        self.put(ws, self.T("cf_loan_in"), L("loan_disb"))
        self.put(ws, self.T("cf_loan_out"), L("loan_princ", "-"))
        self.put(ws, self.T("cf_dividends"), L("div_paid", "-"), key="cf_div")
        b = self.put(ws, self.T("cf_fin_costs"), lambda k, c: f"=-{self.ref('loan_int', c)}-{self.ref('c65_paid', c)}")
        r_fin = self.put(ws, self.T("cf_fin_net"), lambda k, c: f"=SUM({c}{a}:{c}{b})", bold=True, key="cf_fin")
        self.skip(ws)
        r_net = self.put(ws, self.T("cf_net"), lambda k, c: f"={c}{r_op}+{c}{r_inv}+{c}{r_fin}", bold=True, key="cf_net")
        r_cash = ws._next
        self.put(ws, self.T("cf_cash_end"), lambda k, c: (f"={self.col(k - 1)}{r_cash}+" if k else "=") + f"{c}{r_net}",
                 kind="stock", bold=True, key="cash")

    def annual(self, key, title, extra_first=None):
        ws = self.wb.create_sheet(self.S(key))
        ws["A1"] = title
        ws["A1"].font = F_TITLE
        ws["A2"] = self.d["company"]["name"]
        ws.column_dimensions["A"].width = 72
        ws.column_dimensions["B"].width = 10 if key != "kpi" else 2
        cols = ([extra_first] if extra_first else []) + [self.fy_label(i) for i in range(self.nfy)]
        for j, (h1, h2) in enumerate(cols):
            c = CL(3 + j)
            ws.column_dimensions[c].width = 16
            for r, v in ((4, h1), (5, h2)):
                ws[f"{c}{r}"] = v
                ws[f"{c}{r}"].font = F_HDR
                ws[f"{c}{r}"].fill = FILL_HDR
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center", wrap_text=True)
        for c, v in (("A", ""), ("B", self.T("code") if key != "kpi" else "")):
            for r in (4, 5):
                ws[f"{c}{r}"].fill = FILL_HDR
                ws[f"{c}{r}"].font = F_HDR
            ws[f"{c}4"] = v
        ws._next = 6
        ws._ncols = len(cols)
        return ws

    def arow(self, ws, label, code, f, bold=False, indent=0, fmt=EUR, key=None, font=None):
        r = ws._next
        ws._next += 1
        ws.cell(r, 1, ("   " * indent) + label).font = F_BOLD if bold else F_CALC
        ws.cell(r, 2, code).font = Font(name=FONT, color="1F4E5A", bold=bold)
        for j in range(ws._ncols):
            c = CL(3 + j)
            ws[f"{c}{r}"] = f(j, c)
            ws[f"{c}{r}"].number_format = fmt
            ws[f"{c}{r}"].font = font or Font(name=FONT, bold=bold)
            if bold:
                ws[f"{c}{r}"].border = TOP
        if key:
            self.rows[key] = (ws.title, r)
        return r

    def pl(self):
        ws = self.annual("pl", self.T("pl_title"))
        C = self.R("calc")

        def fy(key):
            sheet, r = self.rows[key]
            return lambda j, c: f"=SUMIFS({self.span(r, sheet='calc')},{self.span(3, sheet='calc')},{j + 1})"
        r70 = self.arow(ws, self.T("pl_70"), "70", fy("m70"), indent=1, key="a70")
        r74 = self.arow(ws, self.T("pl_74"), "74", fy("m74"), indent=1, key="a74")
        r60 = self.arow(ws, self.T("pl_60"), "60", fy("m60"), indent=1, key="a60")
        r61 = self.arow(ws, self.T("pl_61"), "61", fy("m61"), indent=1, key="a61")
        r9900 = self.arow(ws, self.T("pl_9900"), "9900", lambda j, c: f"={c}{r70}+{c}{r74}-{c}{r60}-{c}{r61}", bold=True, key="a9900")
        r62 = self.arow(ws, self.T("pl_62"), "62", fy("m62"), indent=1, key="a62")
        r630 = self.arow(ws, self.T("pl_630"), "630", fy("m630"), indent=1, key="a630")
        r640 = self.arow(ws, self.T("pl_640"), "640/8", fy("m640"), indent=1, key="a640")
        r9901 = self.arow(ws, self.T("pl_9901"), "9901", lambda j, c: f"={c}{r9900}-{c}{r62}-{c}{r630}-{c}{r640}", bold=True, key="a9901")
        r65 = self.arow(ws, self.T("pl_65"), "65", fy("m65"), indent=1, key="a65")
        r9903 = self.arow(ws, self.T("pl_9903"), "9903", lambda j, c: f"={c}{r9901}-{c}{r65}", bold=True, key="a9903")
        r_tax = self.arow(ws, self.T("pl_67"), "67/77", lambda j, c: None, indent=1, key="a67")
        r9904 = self.arow(ws, self.T("pl_9904"), "9904", lambda j, c: f"={c}{r9903}-{c}{r_tax}", bold=True, key="a9904")
        self.arow(ws, self.T("pl_9905"), "9905", lambda j, c: f"={c}{r9904}", bold=True, key="a9905")
        ws._next += 1
        ws.cell(ws._next, 1, self.T("tax_calc")).font = F_BOLD
        ws._next += 1
        r_lcf = ws._next
        self.arow(ws, self.T("tax_lcf_start"), "", lambda j, c: "=0" if j == 0 else f"={CL(2 + j)}{r_lcf + 2}")
        r_base = self.arow(ws, self.T("tax_base"), "", lambda j, c: f"=MAX(0,{c}{r9903}-{c}{r_lcf})")
        self.arow(ws, self.T("tax_lcf_end"), "", lambda j, c: f"=MAX(0,{c}{r_lcf}-{c}{r9903})")
        cl_sheet, cl_first, cl_last = self.rows["cost_lines"]
        K = self.R("costs")
        r_rem = self.arow(ws, self.T("tax_max_rem"), "",
                          lambda j, c: f"=_xlfn.MAXIFS({K}${CL(4 + j)}${cl_first}:${CL(4 + j)}${cl_last},{K}$B${cl_first}:$B${cl_last},618)")
        r_el = self.arow(ws, self.T("tax_eligible"), "",
                         # art. 215 WIB 92: the pay condition applies from the fifth taxable period after incorporation
                         lambda j, c: f"=IF(AND(p_UseRed=1,OR({j + 1}<p_TaxRedFrom,{c}{r_rem}>=MIN(p_TaxRedMinRem,{c}{r_base})),{c}{r_base}>0),1,0)",
                         fmt="0")
        self.arow(ws, self.T("tax_amount"), "",
                  lambda j, c: f"=IF({c}{r_el}=1,MIN({c}{r_base},p_TaxRedCap)*p_TaxRed+MAX(0,{c}{r_base}-p_TaxRedCap)*p_TaxStd,{c}{r_base}*p_TaxStd)",
                  bold=True, key="a_tax")
        _, r_amt = self.rows["a_tax"]
        for j in range(self.nfy):
            ws[f"{CL(3 + j)}{r_tax}"] = f"={CL(3 + j)}{r_amt}"
        self.rows["a_eligible"] = (ws.title, r_el)
        self.rows["a_base"] = (ws.title, r_base)
        # monthly tax row on the calculation sheet
        cws = self.wb[self.S("calc")]
        _, rt = self.rows["m_tax"]
        P = self.R("pl")
        for k in range(self.n):
            c = self.col(k)
            cws[f"{c}{rt}"] = f"=IF({c}$4>0,INDEX({P}$C${r_tax}:${CL(2 + self.nfy)}${r_tax},{c}$4),0)"
            cws[f"{c}{rt}"].number_format = EUR
            cws[f"{c}{rt}"].font = F_CALC
        for i in range(self.nfy):
            c = CL(4 + i)
            cws[f"{c}{rt}"] = f"=SUMIFS({self.span(rt)},{self.span(3)},{i + 1})"
            cws[f"{c}{rt}"].number_format = EUR
            cws[f"{c}{rt}"].font = F_BOLD
            cws[f"{c}{rt}"].fill = FILL_TOTAL
        ws.freeze_panes = "C6"

    def balance(self):
        opening = (self.T("bs_opening"), f"{to_date(self.d['company']['start_date']):%d/%m/%Y}")
        ws = self.annual("bs", self.T("bs_title"), extra_first=opening)
        first = self.col(0)

        def fy(key, neg=False, pos=False):
            sheet, r = self.rows[key]
            sh = [k for k, v in SHEETS.items() if v[self.lang] == sheet][0]

            def f(j, c):
                if j == 0:
                    return None
                e = f"SUMIFS({self.span(r, sheet=sh)},{self.span(5, sheet='timeline')},{j})"
                if neg:
                    return f"=MAX(0,-{e})"
                if pos:
                    return f"=MAX(0,{e})"
                return f"={e}"
            return f

        def open_or(fn, open_f):
            return lambda j, c: open_f if j == 0 else fn(j, c)

        ref1 = lambda key: self.ref(key, first)  # noqa: E731
        self.arow(ws, self.T("bs_assets"), "", lambda j, c: None, bold=True)
        self.arow(ws, self.T("bs_20"), "20", lambda j, c: 0, indent=1)
        r21 = self.arow(ws, self.T("bs_21"), "21", open_or(fy("nbv21"), "=0"), indent=2)
        r22 = self.arow(ws, self.T("bs_22"), "22/27", open_or(fy("nbv22"), f"={ref1('inkind')}"), indent=2)
        r28 = self.arow(ws, self.T("bs_28"), "28", open_or(fy("nbv28"), "=0"), indent=2)
        r_fa = self.arow(ws, self.T("bs_2128"), "21/28", lambda j, c: f"={c}{r21}+{c}{r22}+{c}{r28}", bold=True, indent=1)
        r3 = self.arow(ws, self.T("bs_3"), "3", open_or(fy("stock"), "=0"), indent=2)
        r40 = self.arow(ws, self.T("bs_40"), "40", open_or(fy("receivables"), "=0"), indent=2)
        r41 = self.arow(ws, self.T("bs_41"), "41", open_or(fy("vat_bal", neg=True), "=0"), indent=2)
        r54 = self.arow(ws, self.T("bs_54"), "54/58", open_or(fy("cash"), f"={ref1('eq_cash')}+{ref1('loan_disb')}"), indent=2, key="bs_cash")
        r_ca = self.arow(ws, self.T("bs_2958"), "29/58", lambda j, c: f"={c}{r3}+{c}{r40}+{c}{r41}+{c}{r54}", bold=True, indent=1)
        r_ta = self.arow(ws, self.T("bs_total_assets"), "20/58", lambda j, c: f"={c}{r_fa}+{c}{r_ca}", bold=True, key="bs_ta")
        ws._next += 1
        self.arow(ws, self.T("bs_liab"), "", lambda j, c: None, bold=True)
        r10 = self.arow(ws, self.T("bs_10"), "10/11", open_or(fy("cum_contr"), f"={ref1('eq_cash')}+{ref1('inkind')}"), indent=2)
        r14 = self.arow(ws, self.T("bs_14"), "14", open_or(lambda j, c: f"{fy('cum_result')(j, c)}-{fy('cum_div')(j, c)[1:]}", "=0"), indent=2)
        r_eq = self.arow(ws, self.T("bs_1015"), "10/15", lambda j, c: f"={c}{r10}+{c}{r14}", bold=True, indent=1, key="bs_eq")
        bf, bfd = fy("loan_bal_fin"), fy("loan_bal_fnd")
        cf_, cfd = fy("loan_cur_fin"), fy("loan_cur_fnd")
        Fn = self.R("finance")
        _, rbf = self.rows["loan_bal_fin"]
        _, rbd = self.rows["loan_bal_fnd"]
        _, rcf = self.rows["loan_cur_fin"]
        _, rcd = self.rows["loan_cur_fnd"]
        r1704 = self.arow(ws, self.T("bs_1704"), "170/4",
                          lambda j, c: f"={Fn}{first}{rbf}-{Fn}{first}{rcf}" if j == 0 else f"{bf(j, c)}-{cf_(j, c)[1:]}", indent=3)
        r1789 = self.arow(ws, self.T("bs_1789"), "178/9",
                          lambda j, c: f"={Fn}{first}{rbd}-{Fn}{first}{rcd}" if j == 0 else f"{bfd(j, c)}-{cfd(j, c)[1:]}", indent=3)
        r17 = self.arow(ws, self.T("bs_17"), "17", lambda j, c: f"={c}{r1704}+{c}{r1789}", indent=2)
        r42 = self.arow(ws, self.T("bs_42"), "42",
                        lambda j, c: f"={Fn}{first}{rcf}+{Fn}{first}{rcd}" if j == 0 else f"{cf_(j, c)}+{cfd(j, c)[1:]}", indent=3)
        r44 = self.arow(ws, self.T("bs_44"), "44", open_or(fy("payables"), "=0"), indent=3)
        r45 = self.arow(ws, self.T("bs_45"), "45", lambda j, c: "=0" if j == 0 else f"{fy('vat_bal', pos=True)(j, c)}+{fy('tax_payable')(j, c)[1:]}",
                        indent=3)
        r4748 = self.arow(ws, self.T("bs_4748"), "47/48", open_or(fy("div_payable"), "=0"), indent=3)
        r4248 = self.arow(ws, self.T("bs_4248"), "42/48", lambda j, c: f"={c}{r42}+{c}{r44}+{c}{r45}+{c}{r4748}", indent=2)
        r_debt = self.arow(ws, self.T("bs_1749"), "17/49", lambda j, c: f"={c}{r17}+{c}{r4248}", bold=True, indent=1)
        r_tl = self.arow(ws, self.T("bs_total_liab"), "10/49", lambda j, c: f"={c}{r_eq}+{c}{r_debt}", bold=True)
        ws._next += 1
        self.arow(ws, self.T("bs_check"), "", lambda j, c: f"=ROUND({c}{r_ta}-{c}{r_tl},2)", key="bs_diff",
                  font=Font(name=FONT, italic=True, color="808080"))
        ws.freeze_panes = "C6"

    def kpis(self):
        ws = self.annual("kpi", self.T("kpi_title"))
        P, B = self.R("pl"), self.R("bs")

        def pl(key):
            return lambda j, c: f"={P}{c}{self.rows[key][1]}"

        def bs(key):
            return lambda j, c: f"={B}{CL(4 + j)}{self.rows[key][1]}"
        r_rev = self.arow(ws, self.T("pl_70"), "", pl("a70"), font=F_LINK)
        r_gm = self.arow(ws, self.T("pl_9900"), "", pl("a9900"), font=F_LINK)
        self.arow(ws, self.T("kpi_gm_pct"), "", lambda j, c: f"=IF({c}{r_rev}=0,0,{c}{r_gm}/{c}{r_rev})", fmt=PCT)
        r_op = self.arow(ws, self.T("pl_9901"), "", pl("a9901"), font=F_LINK)
        self.arow(ws, self.T("kpi_op_pct"), "", lambda j, c: f"=IF({c}{r_rev}=0,0,{c}{r_op}/{c}{r_rev})", fmt=PCT)
        r_net = self.arow(ws, self.T("pl_9904"), "", pl("a9904"), font=F_LINK)
        self.arow(ws, self.T("kpi_net_pct"), "", lambda j, c: f"=IF({c}{r_rev}=0,0,{c}{r_net}/{c}{r_rev})", fmt=PCT)
        self.arow(ws, self.T("kpi_cash_end"), "", bs("bs_cash"), font=F_LINK)
        _, rc = self.rows["cash"]
        self.arow(ws, self.T("kpi_cash_min"), "",
                  lambda j, c: f"=_xlfn.MINIFS({self.span(rc, sheet='cf')},{self.span(3, sheet='cf')},{j + 1})", key="kpi_min")
        self.arow(ws, self.T("bs_1015"), "", bs("bs_eq"), font=F_LINK)
        _, rf = self.rows["fte"]
        self.arow(ws, self.T("kpi_fte"), "", lambda j, c: f"={self.R('costs')}{CL(4 + j)}{rf}", fmt="0.0", font=F_LINK)
        ws._next += 1
        r = ws._next
        ws.cell(r, 1, self.T("kpi_lowest")).font = F_BOLD
        ws.cell(r, 3, f"=MIN({self.span(rc, sheet='cf')})").number_format = EUR
        ws.cell(r, 4, f"=INDEX({self.span(2, sheet='cf')},MATCH(C{r},{self.span(rc, sheet='cf')},0))").number_format = DATE
        self.rows["kpi_lowest"] = (ws.title, r)
        ch = LineChart()
        ch.title = self.T("chart_cash")
        ch.height, ch.width = 8, 22
        ch.legend = None
        cfws = self.wb[self.S("cf")]
        data = Reference(cfws, min_col=self.mc0, max_col=self.mc0 + self.n - 1, min_row=rc, max_row=rc)
        cats = Reference(cfws, min_col=self.mc0, max_col=self.mc0 + self.n - 1, min_row=2, max_row=2)
        ch.add_data(data, from_rows=True, titles_from_data=False)
        ch.set_categories(cats)
        ch.y_axis.number_format = '€ #,##0'
        ch.x_axis.number_format = "mmm-yy"
        ws.add_chart(ch, f"A{r + 3}")

    def checks(self):
        ws = self.wb.create_sheet(self.S("checks"))
        ws["A1"] = self.T("checks_title")
        ws["A1"].font = F_TITLE
        for c, w in zip("ABCD", (60, 16, 14, 80)):
            ws.column_dimensions[c].width = w
        for i, h in enumerate(("check", "value", "status", "explanation"), 1):
            ws.cell(3, i, self.T(f"ck_{h}")).font = F_HDR
            ws.cell(3, i).fill = FILL_HDR
        ok, err, warn = (f"\"{self.T(x)}\"" for x in ("st_ok", "st_error", "st_warn"))
        B, P, K = self.R("bs"), self.R("pl"), self.R("kpi")
        _, rd = self.rows["bs_diff"]
        _, rc = self.rows["cash"]
        _, rfx = self.rows["fixed"]
        _, rel = self.rows["a_eligible"]
        _, rbase = self.rows["a_base"]
        _, req = self.rows["bs_eq"]
        last_bs = CL(3 + self.nfy)
        last_pl = CL(2 + self.nfy)
        items = [
            ("ck_balance", f"=MAX(ABS({B}C{rd}),ABS({B}{last_bs}{rd}))" if self.nfy == 1 else
             f"=SUMPRODUCT(ABS({B}C{rd}:{last_bs}{rd}))", f"=IF(B{{r}}<0.01,{ok},{err})", EUR2),
            ("ck_cash", f"=MIN({self.span(rc, sheet='cf')})", f"=IF(B{{r}}>=0,{ok},{err})", EUR),
            ("ck_buffer", f"=MIN({self.span(rc, sheet='cf')})-p_CashBuffer*AVERAGE({self.span(rfx, sheet='calc')})",
             f"=IF(B{{r}}>=0,{ok},{warn})", EUR),
            ("ck_reduced", f"=IF(p_UseRed=1,SUMPRODUCT(({P}C{rbase}:{last_pl}{rbase}>0)*(1-{P}C{rel}:{last_pl}{rel})),0)", f"=IF(B{{r}}=0,{ok},{warn})", "0"),
            ("ck_equity", f"=MIN({B}D{req}:{last_bs}{req})", f"=IF(B{{r}}>=0,{ok},{warn})", EUR),
            ("ck_dividend", f"=SUM({self.span(self.rows['div_test'][1], sheet='calc')})", f"=IF(B{{r}}=0,{ok},{err})", "0"),
            # art. 5:4: projections must cover at least two years after incorporation
            ("ck_horizon", f"=EOMONTH(p_FY1End,{12 * (self.nfy - 1)})-(EDATE(p_Start,24)-1)", f"=IF(B{{r}}>=0,{ok},{err})", "0"),
        ]
        r = 4
        for key, val, st, fmt in items:
            ws.cell(r, 1, self.T(key)).font = F_CALC
            ws.cell(r, 2, val).number_format = fmt
            ws.cell(r, 3, st.format(r=r)).font = F_BOLD
            ws.cell(r, 4, self.T(key + "_expl")).font = Font(name=FONT, italic=True, color="666666")
            self.rows[key] = (ws.title, r)
            ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        self.rows["checks"] = (ws.title, 4, r - 1)


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    data = json.loads(Path(sys.argv[1]).read_text())
    Model(data).build().save(sys.argv[2])


if __name__ == "__main__":
    main()
