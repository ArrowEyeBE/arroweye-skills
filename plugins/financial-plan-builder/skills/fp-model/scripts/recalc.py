#!/usr/bin/env python3
"""Recalculate an .xlsx with LibreOffice (headless) and scan for formula errors.

Usage: recalc.py <in.xlsx> [out.xlsx]
Without out.xlsx the input file is replaced by the recalculated version.
Prints JSON: {"status": "success" | "errors_found", "errors": {...}, "formulas": n}
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import plugin_env  # noqa: E402

plugin_env.ensure()

from openpyxl import load_workbook  # noqa: E402

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "Err:")


def soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in (Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
              Path("C:/Program Files/LibreOffice/program/soffice.exe"),
              Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")):
        if p.exists():
            return str(p)
    sys.exit(json.dumps({"status": "error", "message": "LibreOffice not found. Install it from libreoffice.org "
                         "(macOS: brew install --cask libreoffice)."}))


def recalc(src, dst):
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [soffice(), f"-env:UserInstallation=file://{profile}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", tmp, str(src)]
        subprocess.run(cmd, check=True, capture_output=True, timeout=300)
        shutil.move(str(Path(tmp) / Path(src).name), dst)


def scan(path):
    vals = load_workbook(path, data_only=True)
    forms = load_workbook(path)
    errors, n = {}, 0
    for ws in forms.worksheets:
        wv = vals[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v.startswith(ERRORS):
                        errors.setdefault(v, []).append(f"{ws.title}!{c.coordinate}")
    return errors, n


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    src = Path(sys.argv[1]).resolve()
    dst = Path(sys.argv[2]).resolve() if len(sys.argv) == 3 else src
    recalc(src, dst)
    errors, n = scan(dst)
    out = {"status": "errors_found" if errors else "success", "formulas": n,
           "errors": {k: {"count": len(v), "locations": v[:20]} for k, v in errors.items()}}
    print(json.dumps(out, indent=2))
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
