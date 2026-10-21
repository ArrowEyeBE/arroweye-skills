"""Read values from a recalculated model via its hidden _map sheet."""
import json

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL


class ModelValues:
    def __init__(self, path):
        self.wb = load_workbook(path, data_only=True)
        m = self.wb["_map"]
        self.meta, self.rows, self.ranges = {}, {}, {}
        for r in m.iter_rows(values_only=True):
            if r[0] == "meta":
                self.meta[r[1]] = r[2]
            elif r[0] == "row":
                self.rows[r[1]] = (r[2], r[3])
                if len(r) > 4 and r[4] is not None:
                    self.ranges[r[1]] = (r[2], r[3], r[4])  # sheet, first row, last row
        self.nfy = int(self.meta["nfy"])
        self.n = int(self.meta["n"])
        self.mc0 = int(self.meta["first_month_col"])
        self.sheets = json.loads(self.meta["sheets"])
        self.fy_labels = json.loads(self.meta["fy_labels"])

    def cell(self, name, col):
        sheet, r = self.rows[name]
        return self.wb[sheet][f"{col}{r}"].value

    def annual(self, name, opening=False):
        """Values per financial year from an annual sheet (P&L, balance sheet, key figures)."""
        sheet, _ = self.rows[name]
        first = 4 if sheet == self.sheets["bs"] and not opening else 3
        n = self.nfy + (1 if opening else 0)
        return [self.cell(name, CL(first + j)) for j in range(n)]

    def monthly(self, name):
        return [self.cell(name, CL(self.mc0 + k)) for k in range(self.n)]

    def months(self):
        ws = self.wb[self.sheets["cf"]]
        return [ws[f"{CL(self.mc0 + k)}2"].value for k in range(self.n)]

    def fy_totals(self, name):
        """Financial-year total columns (D..) of a monthly sheet."""
        return [self.cell(name, CL(4 + j)) for j in range(self.nfy)]
