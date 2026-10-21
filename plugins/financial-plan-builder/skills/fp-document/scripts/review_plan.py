#!/usr/bin/env python3
"""Mechanical review of a financial plan, the input for the fp-reviewer agent.

Usage: review_plan.py <plan-data.json> <recalculated model.xlsx> [<document.docx>]

Prints JSON:
  facts     what the numbers say: per financial year, streams, staff, investments, financing
  texts     the owner texts with placeholders filled in, per plan-data field
  findings  [{severity: error|warning|info, code, where, message, judge}]
            judge = true: a keyword signal; the reviewer reads the text and confirms or dismisses it.
Exit code 1 if a finding has severity error and judge false.
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "fp-model/scripts"))
from build_document import PH, PlanDocument  # noqa: E402
from doc_i18n import TEXT  # noqa: E402

# Keyword signals per language (regex, case-insensitive). They only point the reviewer to a sentence.
SIGNALS = {
    "hire": {
        "nl": [r"\baanwerv", r"\baangeworven", r"\bin dienst (genomen|treden|komt|komen)", r"\bkomt .{0,40}in dienst",
               r"\bnieuwe? (medewerker|werknemer|collega)", r"\bpersoneelslid"],
        "fr": [r"\bembauch", r"\brecrut", r"\bengag(e|é|ée|ement) (un|une|d)", r"\bnouvel(le)? (collaborat|salari|employ)"],
        "en": [r"\bhir(e|es|ed|ing)\b", r"\brecruit", r"\bnew (employee|staff member|colleague)", r"\bjoins? the team"],
    },
    "loan": {
        "nl": [r"\blening", r"\bkrediet", r"\bbankfinanciering", r"\bgeleend"],
        "fr": [r"\bprêt", r"\bemprunt", r"\bcrédit\b"],
        "en": [r"\bloan", r"\bborrow", r"\bcredit (line|facility)"],
    },
    "no_loan": {
        "nl": [r"\bgeen (enkele )?(lening|krediet|externe financiering|schulden)", r"\bzonder (lening|krediet|externe financiering)",
               r"\b(uitsluitend|volledig|enkel) (met |uit )?eigen middelen"],
        "fr": [r"\b(aucun|pas d'|pas de)\s?(prêt|emprunt|crédit|financement externe)", r"\bsans (prêt|emprunt|financement externe)",
               r"\b(uniquement|entièrement) (par|sur|avec) (des |les )?fonds propres"],
        "en": [r"\bno (loans?|borrowing|external financing|debt)\b", r"\bwithout (a |any )?(loan|borrowing|external financing)",
               r"\b(entirely|only|solely) (from|with|by) (own|the founders'?) (funds|equity)"],
    },
    "grant": {
        "nl": [r"\bsubsidie", r"\bpremie\b", r"\bvlaio\b", r"\bsteunmaatregel"],
        "fr": [r"\bsubside", r"\bprime (de |à l)", r"\bsubvention"],
        "en": [r"\bgrant", r"\bsubsid"],
    },
}

STOPWORDS = {
    "nl": {"de", "het", "een", "en", "van", "voor", "met", "wordt", "zijn", "op", "niet", "ook"},
    "fr": {"le", "la", "les", "des", "et", "du", "pour", "avec", "est", "sont", "une", "pas"},
    "en": {"the", "and", "of", "for", "with", "is", "are", "to", "a", "an", "not", "will"},
}
GOODS_TYPES = {"trade", "hospitality", "production"}
NUM = re.compile(r"(€\s?)?(?<![\w/:.,°])(\d{1,3}(?:[ .,  ]\d{3})+(?:[.,]\d+)?|\d+(?:[.,]\d+)?)(?![\w:°])(\s?(?:%|€|EUR\b|euro\b))?")


def ym(s):
    y, m = str(s)[:7].split("-")
    return int(y), int(m)


def months_between(a, b):
    return (b[0] - a[0]) * 12 + b[1] - a[1]


def parse_number(tok, lang):
    t = tok.replace(" ", " ").replace(" ", " ")
    if lang == "en":
        t = t.replace(",", "").replace(" ", "")
    else:
        t = t.replace(" ", "")
        if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):
            t = t.replace(".", "")
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def numeric_leaves(x, skip=("expected", "activity", "assumptions", "document", "founders", "defaults_used", "expect_status")):
    """Every number in plan-data (inputs), including numbers inside names like 'Huurwaarborg (3 maanden)'."""
    out = []
    if isinstance(x, dict):
        for k, v in x.items():
            if k not in skip:
                out += numeric_leaves(v, ())
    elif isinstance(x, list):
        for v in x:
            out += numeric_leaves(v, ())
    elif isinstance(x, bool):
        pass
    elif isinstance(x, (int, float)):
        out.append(float(x))
    elif isinstance(x, str) and not re.fullmatch(r"\d{4}-\d{2}(-\d{2})?", x):
        out += [float(n) for n in re.findall(r"\d+(?:\.\d+)?", x)]
    return out


def close(a, b):
    return abs(a - b) <= max(0.5, abs(b) * 0.01)


class Review:
    def __init__(self, data, model, docx=None):
        self.d = data
        self.lang = data["meta"]["language"]
        self.pd = PlanDocument(data, model, draft=True)
        self.mv = self.pd.mv
        self.nfy = self.mv.nfy
        self.docx = docx
        self.findings = []
        c = data["company"]
        self.start = ym(c["start_date"])
        fy1 = ym(c["first_fy_end"])
        self.fy_end = [(fy1[0] + i, fy1[1]) for i in range(self.nfy)]

    def add(self, severity, code, where, message, judge=False):
        if any(f["code"] == code and f["where"] == where and f["message"] == message for f in self.findings):
            return
        self.findings.append({"severity": severity, "code": code, "where": where, "message": message, "judge": judge})

    def fy_of(self, month):
        m = ym(month)
        for i, e in enumerate(self.fy_end):
            if months_between(m, e) >= 0:
                return i + 1 if months_between(self.start, m) >= -1 else 0
        return None

    def pay_line(self, c, fy):
        """Director pay of a 618 line that starts or changes in financial year fy (amount + interval, or a schedule)."""
        f = self.pd.f
        sch = {m: v for m, v in (c.get("schedule") or {}).items() if v}
        if sch:
            steps, prev = [], None
            for m in sorted(sch):
                if sch[m] != prev:
                    steps.append((m, sch[m]))
                prev = sch[m]
            here = [f"{f.money(v)} from {m}" for m, v in steps if self.fy_of(m) == fy]
            return f"{c['name']}: " + ", ".join(here) + " (monthly schedule)" if here else None
        if not c.get("start") or self.fy_of(c["start"]) != fy:
            return None
        return (f"{c['name']}: {f.eur(c['amount'])}/{c.get('interval_months', 1)}m from {c['start']}"
                + (f" to {c['end']}" if c.get("end") else ""))

    # ---- facts
    def facts(self):
        mv, f = self.mv, self.pd.f
        ann = {k: mv.annual(r) for k, r in (("revenue", "a70"), ("gross_margin", "a9900"), ("operating_result", "a9901"),
                                             ("net_result", "a9904"), ("cash_end", "bs_cash"), ("equity", "bs_eq"),
                                             ("lowest_cash", "kpi_min"), ("goods", "a60"))}
        fte = mv.fy_totals("fte")
        months = self.pd.fy_months()
        years = []
        for j, (lab, period) in enumerate(mv.fy_labels):
            rev = ann["revenue"][j] or 0
            years.append({
                "fy": j + 1, "label": lab, "period": period, "months": months[j],
                **{k: round(v[j] or 0) for k, v in ann.items()},
                "gm_pct": round((ann["gross_margin"][j] or 0) / rev * 100, 1) if rev else None,
                "fte_end": round(fte[j] or 0, 2),
                "staff_starting": [s["role"] for s in self.d.get("staff", []) if self.fy_of(s["start"]) == j + 1],
                "staff_ending": [s["role"] for s in self.d.get("staff", []) if s.get("end") and self.fy_of(s["end"]) == j + 1],
                "investments": [f"{i['name']} ({f.eur(i['amount'])}, {i['date']})" for i in self.d.get("investments", [])
                                if self.fy_of(i["date"]) == j + 1],
                "director_pay": [self.pay_line(c, j + 1) for c in self.d.get("cost_lines", [])
                                 if c.get("category") == 618 and self.pay_line(c, j + 1)],
            })
        streams = []
        for sid, name, vol_row, rev_row in self.pd.lists.get("streams", []):
            sd = next(x for x in self.d["revenue_streams"] if x["id"] == sid)
            streams.append({"id": sid, "name": name, "unit": sd.get("unit"), "price_excl_vat": sd.get("price_excl_vat"),
                            "price_indexation": sd.get("price_indexation", 0), "payment_lag_months": sd.get("payment_lag_months", 0),
                            "goods_cost_pct": sd.get("goods_cost_pct", 0), "price_schedule": sd.get("price_schedule"),
                            "payment_note": sd.get("payment_note"),
                            "volume_per_fy": [round(v or 0, 1) for v in self.pd.fy_vals("revenue", vol_row)],
                            "revenue_per_fy": [round(v or 0) for v in self.pd.fy_vals("revenue", rev_row)],
                            "first_month_with_volume": next((m for m, v in sorted(sd.get("volumes", {}).items()) if v), None)})
        fin = self.d.get("financing", {})
        return {
            "company": {k: self.d["company"].get(k) for k in ("name", "start_date", "first_fy_end", "business_types", "address")},
            "language": self.lang,
            "financial_years": years,
            "revenue_streams": streams,
            "staff": self.d.get("staff", []),
            "investments": self.d.get("investments", []),
            "cost_lines": [{k: c.get(k) for k in ("name", "category", "amount", "interval_months", "start", "end", "schedule", "source")
                            if c.get(k) is not None} for c in self.d.get("cost_lines", [])],
            "financing": {"equity": fin.get("equity", []), "loans": fin.get("loans", []), "grants": fin.get("grants", []),
                          "dividends": fin.get("dividends", []), "unavailable_equity": fin.get("unavailable_equity", 0)},
            "founders": [{k: x.get(k) for k in ("name", "role", "side_activity")} for x in self.d.get("founders", [])],
            "lowest_cash": self.pd.summary["lowest_cash"],
            "checks": self.pd.summary["checks"],
            "legal_rules": self.pd.summary["legal_rules"],
        }

    # ---- texts
    def text_fields(self):
        """(path, raw text, financial year or None) for every owner text."""
        out = []
        a, act, doc = self.d.get("assumptions") or {}, self.d.get("activity") or {}, self.d.get("document") or {}
        for k in ("purpose", "offer", "business_model", "market", "outlook"):
            if act.get(k):
                out.append((f"activity.{k}", act[k], None))
        for i, t in enumerate(a.get("years") or []):
            out.append((f"assumptions.years[{i}]", t, i + 1))
        for k in ("profitability", "risks"):
            if a.get(k):
                out.append((f"assumptions.{k}", a[k], None))
        if doc.get("financing_note"):
            out.append(("document.financing_note", doc["financing_note"], None))
        return out

    def texts(self):
        return {p: self.pd.resolve(t, p) for p, t, _ in self.text_fields()}

    # ---- checks
    def check_model(self):
        s = self.pd.summary
        for c in s["checks"]:
            if c["status"] in ("error", "warn"):
                self.add("error" if c["status"] == "error" else "warning", f"model_{c['check']}", "model",
                         f"model check {c['check']} is {c['status']} (value {c['value']})")
        lr = s["legal_rules"]
        if lr["stale"]:
            self.add("warning", "rules_stale", "belgian-rules.json",
                     f"Belgian rules last checked {lr['last_checked']} ({lr['age_days']} days ago): run fp-rules-check")
        if lr["to_verify"]:
            self.add("info", "rules_to_verify", "belgian-rules.json", "values marked to verify: " + ", ".join(lr["to_verify"]))

    def check_legal(self):
        act, a, doc = self.d.get("activity") or {}, self.d.get("assumptions") or {}, self.d.get("document") or {}
        for k in ("purpose", "offer", "business_model"):
            if not (act.get(k) or "").strip():
                self.add("error", "text_missing", f"activity.{k}", "required for art. 5:4, §2, 1° (description of the activity)")
        desc = len((act.get("offer") or "") + (act.get("business_model") or ""))
        if 0 < desc < 400:
            self.add("warning", "description_thin", "activity",
                     f"offer + business model are {desc} characters; art. 5:4, §2, 1° asks for a precise description")
        years = a.get("years") or []
        if len(years) != self.nfy:
            self.add("error", "years_count", "assumptions.years", f"{len(years)} year texts for {self.nfy} financial years (art. 5:4, §2, 6°)")
        for i, t in enumerate(years):
            if not PH.search(t) and not re.search(r"\d", t):
                self.add("info", "year_no_figures", f"assumptions.years[{i}]",
                         "no figures or drivers; 6° asks for the hypotheses behind revenue and profitability")
        for k in ("profitability", "risks"):
            if not (a.get(k) or "").strip():
                self.add("error", "text_missing", f"assumptions.{k}", "required (art. 5:4, §2, 6°)")
        if not doc.get("date") or not doc.get("place"):
            self.add("warning", "doc_date_place", "document", "date and place of the plan are needed for the final document")
        elif doc["date"] > self.d["company"]["start_date"]:
            self.add("warning", "doc_after_start", "document.date",
                     f"plan dated {doc['date']}, after the start date {self.d['company']['start_date']}: "
                     "the founders hand the plan to the notary before incorporation (art. 5:4, §1)")
        for i, fo in enumerate(self.d.get("founders", [])):
            for k in ("name", "address"):
                if not fo.get(k):
                    self.add("error", "founder_incomplete", f"founders[{i}].{k}", "missing")
        if not self.d.get("founders"):
            self.add("error", "founder_incomplete", "founders", "no founders listed")

    def check_document(self):
        if not self.docx:
            return
        from docx import Document
        d = Document(self.docx)
        parts = [p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells]
        for s in d.sections:
            parts += [p.text for p in s.header.paragraphs]
        text = "\n".join(parts)
        T = lambda k: TEXT[k][self.lang]  # noqa: E731
        for n in range(1, 8):
            if T("legal_ref").format(n=n) not in text:
                self.add("error", "legal_element_missing", "document", f"art. 5:4, §2, {n}° not found in the document")
        for k in ("h_disclaimer", "h_signature"):
            if T(k) not in text:
                self.add("error", "section_missing", "document", f"section '{T(k)}' missing")
        if set(self.pd.fy_months()) != {12} and T("periodicity_note").split("{")[0] not in text:
            self.add("error", "periodicity_note_missing", "document", "first financial year is not 12 months but the §3 note is missing")
        if T("draft") in text or "[?" in text or T("missing").split("{")[0] in text:
            self.add("error", "draft_document", "document", "this is a draft (draft marker or gaps present); build the final document")

    def signal(self, kind, text):
        hits = []
        for pat in SIGNALS[kind][self.lang]:
            hits += [m for m in re.finditer(pat, text, re.I)]
        return hits

    def sentence(self, text, pos):
        a = max(text.rfind(c, 0, pos) for c in ".\n!?") + 1
        b = min([x for x in (text.find(c, pos) for c in ".\n!?") if x >= 0] or [len(text)])
        return text[a:b + 1].strip()

    NEGATION = {"nl": r"\b(geen|niet|noch|zonder)\b", "fr": r"\b(aucun\w*|pas|ni|sans)\b", "en": r"\b(no|not|nor|without|none)\b"}

    def negated(self, text, pos):
        """True when the sentence part before the signal word negates it ("geen nieuwe ... aanwervingen")."""
        a = max(text.rfind(c, 0, pos) for c in ".\n!?;:") + 1
        return bool(re.search(self.NEGATION[self.lang], text[a:pos], re.I))

    def check_claims(self):
        fin = self.d.get("financing", {})
        staff = self.d.get("staff", [])
        loans, grants = fin.get("loans", []), fin.get("grants", [])
        resolved = self.texts()
        for path, _, fy in self.text_fields():
            if path == "activity.purpose":
                continue  # articles of association: may list activities like lending
            t = resolved[path]
            # hires
            if path not in ("activity.outlook", "assumptions.risks"):
                for m in self.signal("hire", t):
                    if self.negated(t, m.start()):
                        continue
                    if fy is not None:
                        starting = [s["role"] for s in staff if self.fy_of(s["start"]) == fy]
                        if not starting:
                            self.add("error", "hire_not_in_model", path,
                                     f"mentions a hire in FY {fy}, but no staff starts in FY {fy}: \"{self.sentence(t, m.start())}\"", judge=True)
                    elif not staff:
                        self.add("error", "hire_not_in_model", path,
                                 f"mentions a hire, but the plan has no staff: \"{self.sentence(t, m.start())}\"", judge=True)
            # loans
            neg = self.signal("no_loan", t)
            if loans:
                for m in neg:
                    self.add("error", "denies_existing_loan", path,
                             f"says there is no loan, but the plan has {len(loans)}: \"{self.sentence(t, m.start())}\"", judge=True)
            elif path != "assumptions.risks":
                for m in self.signal("loan", t):
                    if not any(n.start() <= m.start() < n.end() for n in neg):
                        self.add("error", "loan_not_in_model", path,
                                 f"mentions a loan, but the plan has none: \"{self.sentence(t, m.start())}\"", judge=True)
            if not grants and path != "assumptions.risks":
                for m in self.signal("grant", t):
                    self.add("error", "grant_not_in_model", path,
                             f"mentions a grant, but the plan has none: \"{self.sentence(t, m.start())}\"", judge=True)
        # what the numbers do that the year texts don't explain
        years = (self.d.get("assumptions") or {}).get("years") or []
        for j in range(min(len(years), self.nfy)):
            t = resolved.get(f"assumptions.years[{j}]", "").lower()
            new = [s for s in staff if self.fy_of(s["start"]) == j + 1]
            if new and not self.signal("hire", t) and not any(s["role"].split()[0].lower() in t for s in new):
                self.add("info", "hire_not_explained", f"assumptions.years[{j}]",
                         "staff starting this year not mentioned: " + ", ".join(s["role"] for s in new))
            inv = [i for i in self.d.get("investments", []) if self.fy_of(i["date"]) == j + 1
                   and months_between(self.start, ym(i["date"])) > 1 and i.get("category", 22) != 28]
            for i in inv:
                if i["name"].split()[0].lower() not in t:
                    self.add("info", "investment_not_explained", f"assumptions.years[{j}]",
                             f"investment during the year not mentioned: {i['name']} ({i['date']})")
        for x in loans:
            if x.get("type") == "founder":
                self.add("info", "founder_loan", "financing.loans",
                         f"founder loan '{x['lender']}' of {self.pd.f.eur(x['amount'])}: confirm every founder agreed to lend it, "
                         "and that a written loan agreement will be made")

    def check_language(self):
        for path, raw, _ in self.text_fields():
            words = re.findall(r"[a-zà-ÿ']+", raw.lower())
            if len(words) < 15:
                continue
            score = {lg: sum(w in sw for w in words) for lg, sw in STOPWORDS.items()}
            best = max(score, key=score.get)
            if best != self.lang and score[best] > 1.5 * score[self.lang]:
                self.add("error", "wrong_language", path, f"text seems to be in '{best}', the plan language is '{self.lang}'")

    def check_numbers(self):
        known = numeric_leaves(self.d)
        known += [v * 100 for v in known if 0 < v < 1]
        known += [v * 30 for v in known if 0 < v <= 12 and v == int(v)]  # payment terms in months → days
        # monthly costs and salaries quoted per year
        known += [c["amount"] * 12 for c in self.d.get("cost_lines", []) if c.get("amount") and c.get("interval_months", 1) == 1]
        known += [s[k] * 12 for s in self.d.get("staff", []) for k in ("gross_monthly", "employer_cost_monthly") if s.get(k)]
        vols = [float(v) for s in self.d.get("revenue_streams", []) for v in s.get("volumes", {}).values() if v]
        outputs = []
        for r in ("a70", "a74", "a60", "a61", "a9900", "a62", "a9901", "a9903", "a67", "a9904", "bs_cash", "bs_eq", "bs_ta"):
            outputs += [v for v in self.mv.annual(r) if v]
        for path, raw, _ in self.text_fields():
            if path == "activity.purpose":
                continue
            text = PH.sub(" ", raw)
            found = []
            for m in NUM.finditer(text):
                v = parse_number(m.group(2), self.lang)
                if v is None:
                    continue
                pct = (m.group(3) or "").strip() == "%"
                money = bool(m.group(1)) or ((m.group(3) or "").strip() not in ("", "%")) or v >= 1000
                if not pct and not money and (v == int(v) and (0 <= v <= 12 or 1990 <= v <= 2100)):
                    continue
                found.append((m.group(0).strip(), v, pct, money))
            plain = [v for _, v, p, _ in found if not p]
            for tok, v, pct, money in found:
                if pct:
                    ok = any(close(v, k) for k in known if k <= 100)
                else:
                    ok = any(close(v, k) for k in known) or any(close(v * o, vol) for o in plain if o != v for vol in vols)
                if money and any(close(v, o) for o in outputs):
                    self.add("warning", "typed_model_figure", path, f"'{tok}' is a model figure typed as text: use a placeholder")
                elif not ok:
                    self.add("warning" if money else "info", "untraced_number", path,
                             f"'{tok}' does not match any input in plan-data" + (": use a placeholder or correct it" if money
                                                                                 else ": confirm with the owner"))

    def check_plausibility(self):
        types = set(self.d["company"].get("business_types", []))
        rev, goods = self.mv.annual("a70"), self.mv.annual("a60")
        gm, net = self.mv.annual("a9900"), self.mv.annual("a9904")
        months = self.pd.fy_months()
        if types & GOODS_TYPES and not any(goods):
            self.add("warning", "no_goods_cost", "revenue_streams",
                     f"business type {', '.join(sorted(types & GOODS_TYPES))} but no cost of goods (60) in any year")
        for j in range(self.nfy):
            if rev[j] and gm[j] < 0:
                self.add("warning", "negative_margin", f"FY {j + 1}", "gross margin (9900) is negative")
        for j in range(1, self.nfy):
            a, b = (rev[j - 1] or 0) / months[j - 1] * 12, (rev[j] or 0) / months[j] * 12
            if a > 0 and b / a > 3:
                self.add("info", "steep_growth", f"FY {j + 1}",
                         f"annualised revenue grows {b / a:.1f}× compared with FY {j}: the year text should explain why")
        if net and all((v or 0) < 0 for v in net):
            self.add("warning", "never_profitable", "assumptions.profitability",
                     "net result negative in every year: explain when and how the company becomes profitable (6°)")
        if not self.pd.has_vat():
            self.add("warning", "no_vat", "revenue_streams, cost_lines",
                     "every VAT rate is 0%: confirm the company will not charge VAT. If it uses the small-business exemption, "
                     "costs must include the VAT it cannot recover (the model does not support the exemption itself)")
        if not any(c.get("category") == 618 for c in self.d.get("cost_lines", [])):
            self.add("info", "no_director_pay", "cost_lines", "no director remuneration in the plan: confirm the founders can live without it")
        for s in self.d.get("revenue_streams", []):
            vols = sorted((m, v) for m, v in s.get("volumes", {}).items())
            jumps = [(vols[k][0], vols[k - 1][1], vols[k][1]) for k in range(1, len(vols))
                     if vols[k - 1][1] and vols[k][1] / vols[k - 1][1] >= 1.5]
            for m, a, b in sorted(jumps, key=lambda x: -x[2] / x[1])[:3]:
                self.add("info", "volume_jump", f"revenue_streams.{s['id']}",
                         f"volume jumps from {a:g} to {b:g} in {m}: check that the year text explains it")

    def run(self):
        self.check_model()
        self.check_legal()
        self.check_document()
        self.check_language()
        self.check_claims()
        self.check_numbers()
        self.check_plausibility()
        order = {"error": 0, "warning": 1, "info": 2}
        self.findings.sort(key=lambda x: order[x["severity"]])
        return {"facts": self.facts(), "texts": self.texts(), "findings": self.findings,
                "reviewed_on": date.today().isoformat()}


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    if len(sys.argv) not in (3, 4):
        sys.exit(__doc__)
    data = json.loads(Path(sys.argv[1]).read_text())
    out = Review(data, sys.argv[2], sys.argv[3] if len(sys.argv) == 4 else None).run()
    print(json.dumps(out, indent=2, ensure_ascii=False, default=str))
    sys.exit(1 if any(f["severity"] == "error" and not f["judge"] for f in out["findings"]) else 0)


if __name__ == "__main__":
    main()
