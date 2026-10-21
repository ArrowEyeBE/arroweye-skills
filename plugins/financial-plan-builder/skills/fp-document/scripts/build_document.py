#!/usr/bin/env python3
"""Build the financial-plan Word document from plan-data.json and the recalculated model.

Usage: build_document.py <plan-data.json> <model.xlsx> <out.docx> [--draft]

<model.xlsx> must be recalculated (recalc.py). Every number in the document comes from it:
tables are read from the workbook, and numbers inside the owner's text are written as
placeholders ({{revenue.1}}, see references/document-text.md) that are filled from it.
Without --draft the build fails when text is missing, a placeholder is unknown or a check
of the model is in error. With --draft, gaps are marked in the document instead.
"""
import argparse
import calendar
import json
import re
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
MODEL_SCRIPTS = HERE.parent.parent / "fp-model" / "scripts"
sys.path.insert(0, str(MODEL_SCRIPTS))
sys.path.insert(0, str(HERE))
import plugin_env  # noqa: E402

plugin_env.ensure()

from check_model import summarise  # noqa: E402
from doc_i18n import MONTHS, MONTHS_SHORT, TEXT  # noqa: E402
from i18n import LABELS  # noqa: E402
from model_io import ModelValues  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

RULES = json.loads((MODEL_SCRIPTS.parent / "references/belgian-rules.json").read_text())
FONT = "Arial"
TEAL = "1F4E5A"
LIGHT = "D9E2E5"
GREY = "F2F2F2"
TEXT_WIDTH = 16.6  # cm, A4 portrait with 2.2 cm margins
LANG_TAG = {"nl": "nl-BE", "fr": "fr-BE", "en": "en-GB"}
ALIAS = {"revenue": "a70", "grants": "a74", "goods": "a60", "services": "a61", "gross_margin": "a9900", "staff_cost": "a62",
         "depreciation": "a630", "operating_result": "a9901", "financial_costs": "a65", "profit_before_tax": "a9903",
         "corporate_tax": "a67", "net_result": "a9904", "cash_end": "bs_cash", "equity": "bs_eq", "total_assets": "bs_ta",
         "lowest_cash_fy": "kpi_min"}
RULE_LABEL = {"reduced_rate_min_director_pay": "tax_red_minrem", "employer_ssc_rate": "ssc", "corporate_tax_standard": "tax_std",
              "corporate_tax_reduced": "tax_red", "corporate_tax_reduced_bracket": "tax_red_cap", "wage_indexation": "wage_idx",
              "salary_months_per_year": "sal_months"}
PH = re.compile(r"\{\{\s*([A-Za-z0-9_.\-]+)(?::(abs|num))?\s*\}\}")
# a unit written next to a placeholder that already carries it ("{{revenue.1}} euro", "{{gm_pct.1}}%")
UNIT_AFTER = re.compile(r"\s?(euro\b|eur\b|€|%)", re.I)
UNIT_BEFORE = re.compile(r"(€|\beur|\beuro)\s?$", re.I)
LEGAL_FORM = re.compile(r"[\s,]+(BV|B\.V\.|SRL|S\.R\.L\.|BVBA|SPRL)\.?\s*$", re.I)


def company_name(c):
    """company.name without a trailing legal form: the document adds it."""
    return LEGAL_FORM.sub("", c["name"].strip())


class BuildError(Exception):
    pass


# ---- formatting ---------------------------------------------------------------
class Fmt:
    def __init__(self, lang):
        self.lang = lang
        self.th, self.dec = {"nl": (".", ","), "fr": (" ", ","), "en": (",", ".")}[lang]

    def num(self, v, d=0, dash=True):
        if v is None or v == "":
            return ""
        if dash and abs(v) < 0.5 * 10 ** -d:
            return "-"
        s = f"{abs(v):,.{d}f}".replace(",", "\x00").replace(".", self.dec).replace("\x00", self.th)
        return ("-" if v < 0 else "") + s

    def eur(self, v, d=0):
        s = self.num(abs(v), d, dash=False)
        sign = "-" if v < 0 and round(abs(v), d) > 0 else ""
        return {"nl": f"{sign}€ {s}", "fr": f"{sign}{s} €", "en": f"{sign}€{s}"}[self.lang]

    def money(self, v):
        return self.eur(v, 0 if abs(v - round(v)) < 0.005 else 2)

    def pct(self, v, d=1):
        s = self.num(v * 100, d, dash=False)
        return f"{s} %" if self.lang == "fr" else f"{s}%"

    def date(self, d):
        d = to_date(d)
        day = "1er" if self.lang == "fr" and d.day == 1 else d.day
        return f"{day} {MONTHS[self.lang][d.month - 1]} {d.year}"

    def month(self, d):
        d = to_date(d)
        return f"{MONTHS[self.lang][d.month - 1]} {d.year}"

    def month_short(self, d):
        d = to_date(d)
        return f"{MONTHS_SHORT[self.lang][d.month - 1]} {d.year % 100:02d}"


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v)
    return date(int(s[:4]), int(s[5:7]), int(s[8:10]) if len(s) >= 10 else 1)


# ---- docx helpers -------------------------------------------------------------
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_border(cell, edge, sz=6, color="808080"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:color"), color)
    b.append(e)


def table_borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge, val, color in (("top", "single", TEAL), ("bottom", "single", TEAL), ("insideH", "single", "D9D9D9"),
                             ("left", "nil", None), ("right", "nil", None), ("insideV", "nil", None)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), val)
        if color:
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def fix_widths(table, widths):
    """Fixed layout with explicit grid and cell widths, so Word and LibreOffice both respect them."""
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))
    for i, w in enumerate(widths):
        table.columns[i].width = Cm(w)
        for cell in table.columns[i].cells:
            cell.width = Cm(w)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        e = OxmlElement(tag)
        e.set(qn("w:val"), "true")
        trPr.append(e)


def no_split(row):
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(e)


def add_field(par, instr, placeholder=""):
    def fld(t):
        r = par.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)
    fld("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    fld("separate")
    par.add_run(placeholder)
    fld("end")


def add_hyperlink(par, url, text, size=9):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(int(size * 2)))):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        rPr.append(e)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    h.append(r)
    par._p.append(h)


# ---- document -----------------------------------------------------------------
class PlanDocument:
    def __init__(self, data, model_path, draft=False):
        self.d = data
        self.lang = data["meta"]["language"]
        self.f = Fmt(self.lang)
        self.mv = ModelValues(model_path)
        if self.mv.meta["lang"] != self.lang:
            raise BuildError(f"model language {self.mv.meta['lang']} differs from plan-data language {self.lang}")
        self.summary = summarise(model_path)
        self.lists = json.loads(self.mv.meta.get("lists") or "{}")
        self.nfy = self.mv.nfy
        self.draft = draft
        self.problems = []
        self.doc = Document()
        self.h1_no = 0
        self.h2_no = 0
        self.toc = []
        self.status = {c["check"]: c["status"] for c in self.summary["checks"]}
        if not draft:
            errs = [c for c, s in self.status.items() if s == "error"]
            if errs:
                raise BuildError(f"model checks in error: {', '.join(errs)}; resolve them first or build with --draft")

    def T(self, key, **kw):
        s = TEXT[key][self.lang]
        return s.format(**kw) if kw else s

    def L(self, key):
        return LABELS[key][self.lang]

    # ---- model access
    def ws(self, key):
        return self.mv.wb[self.mv.sheets[key]]

    def param(self, name):
        dn = self.mv.wb.defined_names[name]
        sheet, coord = next(dn.destinations)
        return self.mv.wb[sheet][coord.replace("$", "")].value

    def fy_vals(self, sheet_key, row):
        ws = self.ws(sheet_key)
        return [ws.cell(row, 4 + j).value or 0 for j in range(self.nfy)]

    def fy_months(self):
        ws = self.ws("cf")
        out = [0] * self.nfy
        for k in range(self.mv.n):
            fy = ws.cell(3, self.mv.mc0 + k).value
            if fy:
                out[int(fy) - 1] += 1
        return out

    def months(self, n):
        """'1 maand', '3 maanden', '1,5 maanden'."""
        txt = self.f.num(n, 1, dash=False).rstrip("0").rstrip(",.") if n != int(n) else str(int(n))
        return self.T("month_1" if n == 1 else "months_n", n=txt)

    def has_vat(self):
        """False when every VAT rate in the plan is 0 (then cash flows equal the amounts excl. VAT)."""
        d = self.d
        rates = [s.get(k, s.get("vat_rate", 0.21)) for s in d.get("revenue_streams", [])
                 for k in ("vat_rate", "purchase_vat_rate", "other_var_vat_rate")]
        rates += [x.get("vat_rate", 0.21) for k in ("cost_lines", "investments") for x in d.get(k, [])]
        return any(rates)

    # ---- placeholders
    def resolve(self, text, where):
        def rep(m):
            key, mod = m.group(1), m.group(2)
            try:
                v = self.placeholder(key, mod)
            except (KeyError, IndexError, ValueError, ZeroDivisionError):
                self.problem(f"unknown placeholder {{{{{key}}}}} in {where}")
                return f"[?{key}?]"
            if mod != "num" and ("€" in v or "%" in v):
                after, before = UNIT_AFTER.match(text, m.end()), UNIT_BEFORE.search(text, 0, m.start())
                if after or before:
                    unit = (after or before).group(1)
                    self.problem(f"{where}: '{unit}' next to {{{{{key}}}}}, which already includes the unit (gives '{v} {unit}'); "
                                 f"remove '{unit}' or use {{{{{key}:num}}}}")
            return v
        return PH.sub(rep, text)

    def placeholder(self, key, mod=None):
        """Formatted value of a placeholder. mod 'abs': without the minus sign; 'num': bare number, no € or % sign."""
        v = self.placeholder_value(key)
        if isinstance(v, str):
            return v
        num, kind = v
        if mod == "abs":
            num = abs(num)
        if kind == "pct":
            return self.f.num(num * 100, 1, dash=False) if mod == "num" else self.f.pct(num)
        if kind == "num":
            return self.f.num(num, 1 if num != int(num) else 0, dash=False)
        if mod == "num":
            return self.f.num(num, 0 if abs(num - round(num)) < 0.005 else 2, dash=False)
        return self.f.money(num) if kind == "money" else self.f.eur(num)

    def placeholder_value(self, key):
        """(number, kind) with kind eur / money / pct / num, or an already formatted text."""
        parts = key.split(".")
        f, mv = self.f, self.mv
        fin = self.d.get("financing", {})
        totals = {
            "equity_cash_total": sum(e["amount"] for e in fin.get("equity", [])),
            "inkind_total": sum(i["amount"] for i in self.d.get("investments", []) if i.get("in_kind")),
            "loans_total": sum(x["amount"] for x in fin.get("loans", [])),
            "grants_total": sum(g["amount"] for g in fin.get("grants", [])),
            "investments_total": sum(i["amount"] for i in self.d.get("investments", [])),
        }
        totals["financing_total"] = totals["equity_cash_total"] + totals["inkind_total"] + totals["loans_total"] + totals["grants_total"]
        if len(parts) == 1:
            if key in totals:
                return totals[key], "eur"
            if key == "lowest_cash":
                return self.summary["lowest_cash"]["amount"], "eur"
            if key == "lowest_cash_month":
                return f.month(self.summary["lowest_cash"]["month"])
            if key == "start_date":
                return f.date(self.d["company"]["start_date"])
            raise KeyError(key)
        if parts[0] == "fy" and len(parts) == 2:
            lab, period = mv.fy_labels[int(parts[1]) - 1]
            return f"{lab} ({period})"
        if parts[0] == "stream" and len(parts) in (3, 4):
            s = next(x for x in self.lists.get("streams", []) if x[0] == parts[1])
            sd = next(x for x in self.d["revenue_streams"] if x["id"] == parts[1])
            if parts[2] == "price" and len(parts) == 3:
                return sd["price_excl_vat"], "money"
            j = int(parts[3]) - 1
            if j < 0:
                raise IndexError(key)
            if parts[2] == "price":  # price in the last month of that financial year (follows a price_schedule)
                ws, k = self.ws("revenue"), self.fy_last_month(j)
                r = next(r for r in range(s[2], s[3]) if ws.cell(r, 1).value == self.L("price_excl"))
                return ws.cell(r, self.mv.mc0 + k).value, "money"
            if parts[2] == "volume":
                return self.fy_vals("revenue", s[2])[j], "num"
            if parts[2] == "revenue":
                return self.fy_vals("revenue", s[3])[j], "eur"
            raise KeyError(key)
        if len(parts) != 2:
            raise KeyError(key)
        name, j = parts[0], int(parts[1]) - 1
        if j < 0:
            raise IndexError(key)
        if name in ("gm_pct", "op_pct", "net_pct"):
            num = {"gm_pct": "a9900", "op_pct": "a9901", "net_pct": "a9904"}[name]
            return mv.annual(num)[j] / mv.annual("a70")[j], "pct"
        if name == "fte":
            return mv.fy_totals("fte")[j], "num"
        row = ALIAS.get(name, name)
        if row not in mv.rows:
            raise KeyError(key)
        return mv.annual(row)[j], "eur"

    def fy_last_month(self, j):
        """Index of the last month of financial year j (0-based) in the monthly sheets."""
        ws = self.ws("cf")
        return max(k for k in range(self.mv.n) if ws.cell(3, self.mv.mc0 + k).value == j + 1)

    def problem(self, msg):
        self.problems.append(msg)

    # ---- styles and page setup
    def setup(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.2))
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
        st = doc.styles
        normal = st["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        rpr = normal.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), FONT)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), LANG_TAG[self.lang])
        rpr.append(lang)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, before, after in (("Heading 1", 16, 0, 10), ("Heading 2", 12, 12, 4), ("Heading 3", 10.5, 8, 2), ("Title", 28, 0, 6)):
            s = st[name]
            s.font.name = FONT
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.italic = False
            s.font.color.rgb = RGBColor.from_string(TEAL)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
            r = s.element.get_or_add_rPr()
            fn = r.find(qn("w:rFonts"))
            if fn is None:
                fn = OxmlElement("w:rFonts")
                r.append(fn)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fn.set(qn(a), FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
                if fn.get(qn(a)) is not None:
                    del fn.attrib[qn(a)]
        st["List Bullet"].font.name = FONT
        c = self.d["company"]
        doc.core_properties.title = f"{self.T('doc_title')} {company_name(c)} {self.T('legal_form_short')}"
        doc.core_properties.author = f"{company_name(c)} {self.T('legal_form_short')}"
        doc.core_properties.language = LANG_TAG[self.lang]
        # ask Word to refresh the table of contents and page numbers on opening
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        doc.settings.element.append(upd)
        self.footer(sec)

    def footer(self, sec):
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        c = self.d["company"]
        r = p.add_run(f"{self.T('doc_title')} {company_name(c)} {self.T('legal_form_short')}  ·  {self.T('page')} ")
        r.font.size = Pt(8)
        add_field(p, "PAGE", "1")
        p.add_run(f" {self.T('of')} ").font.size = Pt(8)
        add_field(p, "NUMPAGES", "1")
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if self.draft:
            hp = sec.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hr = hp.add_run(self.T("draft"))
            hr.bold = True
            hr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # ---- text helpers
    def para(self, text="", size=None, italic=False, bold=False, color=None, align=None, after=None):
        p = self.doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.italic, r.bold = italic, bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if align:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def h1(self, key, legal=None, new_page=True):
        if new_page:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            # drop the empty paragraph's spacing
            self.doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
        self.h1_no += 1
        self.h2_no = 0
        self.doc.add_heading(f"{self.h1_no}. {self.T(key)}", level=1)
        self.toc.append((1, f"{self.h1_no}. {self.T(key)}"))
        if legal:
            self.para(self.T("legal_ref", n=legal), size=8, italic=True, color="666666", after=8)

    def h2(self, key_or_text, raw=False):
        self.h2_no += 1
        t = key_or_text if raw else self.T(key_or_text)
        self.doc.add_heading(f"{self.h1_no}.{self.h2_no} {t}", level=2)
        self.toc.append((2, f"{self.h1_no}.{self.h2_no} {t}"))

    def prose(self, text, what, required=True):
        """Owner text: paragraphs split by blank lines, '- ' lines become bullets, **bold** allowed."""
        if not text or not str(text).strip():
            if required:
                self.problem(f"missing text: {what}")
                p = self.para()
                r = p.add_run(self.T("missing", what=what))
                r.font.highlight_color = 7  # yellow
            return
        text = self.resolve(str(text), what)
        for block in re.split(r"\n\s*\n", text.strip()):
            lines = block.split("\n")
            if all(ln.strip().startswith(("- ", "* ")) for ln in lines):
                for ln in lines:
                    self.rich(self.doc.add_paragraph(style="List Bullet"), ln.strip()[2:])
            else:
                self.rich(self.doc.add_paragraph(), " ".join(ln.strip() for ln in lines))

    @staticmethod
    def rich(p, text):
        for i, part in enumerate(re.split(r"\*\*", text)):
            if part:
                p.add_run(part).bold = i % 2 == 1

    # ---- tables
    def table(self, headers, rows, widths, num_from=1, size=8.5, header_rows=None):
        """rows: list of dicts {cells: [...], bold, section, total, indent} or None for a spacer."""
        rows = [r for r in rows if r is not None]
        t = self.doc.add_table(rows=0, cols=len(widths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_borders(t)
        for hdr in header_rows or [headers]:
            row = t.add_row()
            repeat_header(row)
            for i, h in enumerate(hdr):
                c = row.cells[i]
                shade(c, TEAL)
                p = c.paragraphs[0]
                r = p.add_run(h or "")
                r.bold = True
                r.font.size = Pt(size)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i >= num_from else WD_ALIGN_PARAGRAPH.LEFT
        for spec in rows:
            row = t.add_row()
            no_split(row)
            for i, val in enumerate(spec["cells"]):
                c = row.cells[i]
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                txt = val if isinstance(val, str) else ""
                if i == 0 and spec.get("indent"):
                    p.paragraph_format.left_indent = Cm(0.35 * spec["indent"])
                r = p.add_run(txt)
                r.font.size = Pt(size)
                r.bold = bool(spec.get("bold") or spec.get("section"))
                if i >= num_from:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if spec.get("section"):
                    shade(c, LIGHT)
                elif spec.get("total"):
                    shade(c, GREY)
                    cell_border(c, "top", 6, TEAL)
        fix_widths(t, widths)
        self.para(after=4)
        return t

    def money_cells(self, values):
        return [self.f.num(v) for v in values]

    def fy_headers(self, first="", extra=None):
        top = [first] + ([extra[0]] if extra else []) + [lab for lab, _ in self.mv.fy_labels]
        bottom = [self.T("in_eur")] + ([extra[1]] if extra else []) + [period.replace(" – ", " –\n") for _, period in self.mv.fy_labels]
        return top, bottom

    def statement_rows(self, sheet_key, first, last, ncols):
        ws = self.ws(sheet_key)
        out = []
        for r in range(first, last + 1):
            a = ws.cell(r, 1)
            if a.value is None:
                continue
            raw = str(a.value)
            indent = (len(raw) - len(raw.lstrip(" "))) // 3
            vals = [ws.cell(r, 3 + j).value for j in range(ncols)]
            bold = bool(a.font and a.font.b)
            section = all(v is None for v in vals)
            out.append({"label": raw.strip(), "code": ws.cell(r, 2).value or "", "values": vals,
                        "bold": bold, "section": section, "indent": indent})
        return out

    def statement_table(self, rows, extra=None):
        top, bottom = self.fy_headers(extra=extra)
        n = len(rows[0]["values"])
        label_w, code_w = (7.6, 1.3) if n <= 3 else (6.0, 1.2)
        widths = [label_w, code_w] + [(TEXT_WIDTH - label_w - code_w) / n] * n
        hdr1 = [top[0], self.T("c_code")] + top[1:]
        hdr2 = [bottom[0], ""] + bottom[1:]
        specs = []
        for r in rows:
            cells = [r["label"], str(r["code"])] + ([""] * n if r["section"] else self.money_cells([v or 0 for v in r["values"]]))
            specs.append({"cells": cells, "bold": r["bold"], "section": r["section"] and r["indent"] == 0,
                          "total": r["bold"] and not r["section"], "indent": r["indent"]})
        self.table(None, specs, widths, num_from=1, header_rows=[hdr1, hdr2])

    # ---- sections
    def cover(self):
        c = self.d["company"]
        for _ in range(6):
            self.para()
        self.doc.add_paragraph(self.T("doc_title"), style="Title")
        p = self.para(f"{company_name(c)} {self.T('legal_form_short')}", size=20, bold=True, color="333333", after=4)
        self.para(self.T("in_formation"), size=12, color="666666", after=30)
        self.para(self.T("cover_legal"), size=10, italic=True, color="666666", after=30)
        self.para(self.T("cover_start", date=self.f.date(c["start_date"])), size=10)
        self.para(self.T("cover_date", date=self.f.date(self.doc_date())), size=10)
        if self.draft:
            self.para(self.T("draft"), size=12, bold=True, color="C00000")
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.doc.add_heading(self.T("toc"), level=1)
        self.toc_par = self.doc.add_paragraph()

    def fill_toc(self):
        """TOC field whose cached result already lists the chapters, so it reads well before Word refreshes it."""
        p0 = self.toc_par
        entries = self.toc or [(1, self.T("toc_hint"))]

        def fld(par, t):
            r = par.add_run()
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), t)
            r._r.append(e)
        fld(p0, "begin")
        r = p0.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = 'TOC \\o "1-2" \\h \\z \\u'
        r._r.append(it)
        fld(p0, "separate")
        pars = [p0]
        anchor = p0
        for _ in entries[1:]:
            new = OxmlElement("w:p")
            anchor._p.addnext(new)
            from docx.text.paragraph import Paragraph
            anchor = Paragraph(new, p0._parent)
            pars.append(anchor)
        for par, (level, text) in zip(pars, entries):
            run = par.add_run(text)
            run.bold = level == 1
            par.paragraph_format.left_indent = Cm(0.6 * (level - 1))
            par.paragraph_format.space_after = Pt(2 if level == 2 else 4)
        fld(pars[-1], "end")

    def doc_date(self):
        return (self.d.get("document") or {}).get("date") or date.today().isoformat()

    def company(self):
        c = self.d["company"]
        self.h1("h_company")
        months = self.fy_months()
        fys = "\n".join(self.T("fy_line", label=lab, period=period, months=months[i]) for i, (lab, period) in enumerate(self.mv.fy_labels))
        rows = [
            (self.T("f_name"), f"{company_name(c)} {self.T('legal_form_short')}"),
            (self.T("f_legal_form"), self.T("legal_form_long")),
            (self.T("f_address"), c.get("address") or self.missing_inline("company.address")),
            (self.T("f_enterprise_nr"), c.get("enterprise_number") or self.T("f_enterprise_tbd")),
            (self.T("f_start"), self.f.date(c["start_date"])),
            (self.T("f_fys"), fys),
            (self.T("f_vat"), self.T("vat_monthly" if int(self.param("p_VATPeriod")) == 1 else "vat_quarterly")),
        ]
        self.kv_table(rows)

    def missing_inline(self, what):
        self.problem(f"missing: {what}")
        return self.T("missing", what=what)

    def kv_table(self, rows, widths=(5.0, 11.6)):
        t = self.doc.add_table(rows=0, cols=2)
        t.autofit = False
        table_borders(t)
        for k, v in rows:
            row = t.add_row()
            no_split(row)
            for i, txt in enumerate((k, v)):
                c = row.cells[i]
                c.width = Cm(widths[i])
                lines = str(txt).split("\n")
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                for j, ln in enumerate(lines):
                    r = p.add_run(ln)
                    r.bold = i == 0
                    r.font.size = Pt(9.5)
                    if j < len(lines) - 1:
                        r.add_break()
            shade(row.cells[0], GREY)
        fix_widths(t, widths)
        self.para(after=4)

    def founders(self):
        self.h1("h_founders", new_page=False)
        fs = self.d.get("founders") or []
        if not fs:
            self.prose("", "founders")
            return
        self.para(self.T("founders_intro"))
        cols = [("name", "c_name", 4.0), ("role", "c_role", 3.2), ("address", "c_address", 5.0),
                ("birth_year", "c_birth_year", 2.0), ("side_activity", "c_side", 2.4)]
        cols = [c for c in cols if any(f.get(c[0]) not in (None, "") for f in fs)]
        scale = TEXT_WIDTH / sum(w for _, _, w in cols)

        def val(f, k):
            v = f.get(k)
            if isinstance(v, bool):
                return self.T("yes" if v else "no")
            return "" if v is None else str(v)
        self.table([self.T(h) for _, h, _ in cols], [{"cells": [val(f, k) for k, _, _ in cols]} for f in fs],
                   [w * scale for _, _, w in cols], num_from=99, size=9)

    def activity(self):
        a = self.d.get("activity") or {}
        self.h1("h_activity", legal=1)
        for key, req in (("purpose", True), ("offer", True), ("business_model", True), ("market", False), ("outlook", False)):
            if a.get(key) or req:
                self.h2(f"h_{key}")
                self.prose(a.get(key), f"activity.{key}", required=req)
                if key == "purpose" and a.get("purpose_status"):
                    self.para(self.T("purpose_status", status=a["purpose_status"].strip().rstrip(".")), size=9, italic=True, color="555555")

    def financing(self):
        fin = self.d.get("financing", {})
        f = self.f
        self.h1("h_financing", legal=2)
        self.para(self.T("financing_intro"))
        self.h2("h_fin_overview")
        rows, total = [], 0
        inkind = [i for i in self.d.get("investments", []) if i.get("in_kind")]
        for e in fin.get("equity", []):
            rows.append([e["name"], self.T("t_equity_cash"), f.month(e["date"]), f.num(e["amount"])])
            total += e["amount"]
        for i in inkind:
            rows.append([i["name"], self.T("t_equity_kind"), f.month(i["date"]), f.num(i["amount"])])
            total += i["amount"]
        for x in fin.get("loans", []):
            rows.append([x["lender"], self.T(f"t_loan_{x.get('type', 'other')}"), f.month(x["date"]), f.num(x["amount"])])
            total += x["amount"]
        for g in fin.get("grants", []):
            rows.append([g["name"], self.T("t_grant"), f.month(g["date"]), f.num(g["amount"])])
            total += g["amount"]
        specs = [{"cells": r} for r in rows] + [{"cells": [self.T("total"), "", "", f.num(total)], "total": True, "bold": True}]
        self.table([self.T("c_source"), self.T("c_type"), self.T("c_date"), self.T("c_amount")], specs, [7.0, 4.0, 3.0, 2.6], num_from=3)
        notes = []
        confirmed = set(fin.get("confirmed_none") or [])
        for k, empty in (("loans", not fin.get("loans")), ("grants", not fin.get("grants")), ("in_kind", not inkind)):
            if empty and k not in confirmed:
                self.problem(f"financing.confirmed_none: the document would state there are no {k}, but the owner has not confirmed it")
        if not fin.get("loans"):
            notes.append(self.T("no_loans"))
        if not inkind:
            notes.append(self.T("no_inkind"))
        if not fin.get("grants"):
            notes.append(self.T("no_grants"))
        notes.append(self.T("no_other_sources"))
        self.para(" ".join(notes))
        if inkind:
            self.para(self.T("inkind_note"), size=9, italic=True, color="555555")
        if fin.get("loans"):
            self.h2("h_loans")
            lrows = []
            for x in fin["loans"]:
                lrows.append({"cells": [x["lender"], f.num(x["amount"]), f.pct(x.get("rate", 0), 2),
                                        self.months(x.get("term_months", 0)),
                                        self.months(x["grace_months"]) if x.get("grace_months") else self.T("none"),
                                        x.get("guarantees") or self.T("none")]})
            self.table([self.T("c_lender"), self.T("c_amount"), self.T("c_rate"), self.T("c_term"), self.T("c_grace"), self.T("c_guarantees")],
                       lrows, [4.2, 2.2, 1.6, 2.0, 2.2, 4.4], num_from=1)
            self.para(self.T("loan_method"), size=9, italic=True, color="555555")
        note = (self.d.get("document") or {}).get("financing_note")
        if note:
            self.prose(note, "document.financing_note", required=False)
        self.h2("h_sufficiency")
        low = self.summary["lowest_cash"]
        key = "suff_neg" if low["amount"] < 0 else "suff_ok"
        self.para(self.T(key, amount=f.eur(low["amount"]), month=f.month(low["month"])))
        if self.status.get("ck_buffer") == "warn" and low["amount"] >= 0:
            self.para(self.T("suff_buffer_low", months=self.months(self.param("p_CashBuffer"))))
        eq = self.mv.annual("bs_eq")
        self.para(self.T("suff_equity", eq1=f.eur(eq[0]), eqn=f.eur(eq[-1])))
        self.dividends()

    def dividends(self):
        f = self.f
        divs = (self.d.get("financing") or {}).get("dividends") or []
        self.h2("h_dividends")
        if not divs:
            self.para(self.T("div_none"))
            return
        if self.status.get("ck_dividend") == "error":
            self.problem("ck_dividend: a dividend would bring equity below zero or below the unavailable equity (art. 5:142 WVV)")
        self.para(self.T("div_intro"))
        cash, months = self.mv.monthly("cash"), [to_date(m) for m in self.mv.months()]
        rows, short = [], False
        for x in divs:
            fy = int(x["fy"])
            decided = min(to_date(x["date"]), to_date(self.fy_end_date(fy - 1)))
            k = next((i for i, m in enumerate(months) if (m.year, m.month) == (to_date(x["date"]).year, to_date(x["date"]).month)), None)
            low = ""
            if k is not None:
                window = cash[k:k + 12]
                low = f.num(min(v or 0 for v in window)) + ("*" if len(window) < 12 else "")
                short |= len(window) < 12
            lab, period = self.mv.fy_labels[fy - 1]
            rows.append({"cells": [f"{lab} ({period})" + (f"\n{x['name']}" if x.get("name") else ""), f.month(decided),
                                   f.month(x["date"]), f.num(x["amount"]), low]})
        self.table([self.T("c_div_from"), self.T("c_decided"), self.T("c_paid"), self.T("c_amount"), self.T("c_cash_after")],
                   rows, [5.4, 2.8, 2.8, 2.4, 3.2], num_from=3, size=8.5)
        if short:
            self.para(self.T("div_window_note"), size=9, italic=True, color="555555", after=2)
        self.para(self.T("div_tests"), size=9, italic=True, color="555555")

    def fy_end_date(self, j):
        period = self.mv.fy_labels[j][1]
        d, m, y = period.split(" – ")[1].split("/")
        return f"{y}-{m}-{d}"

    def balance(self):
        self.h1("h_balance", legal=3)
        c = self.d["company"]
        txt = self.T("balance_intro", start=self.f.date(c["start_date"]))
        months = self.fy_months()
        if months[0] != 12:
            txt += " " + self.T("first_fy_note", months=months[0], period=self.mv.fy_labels[0][1])
        self.para(txt)
        self.periodicity_note()
        _, first = self.mv.rows["bs_ta"]
        _, diff = self.mv.rows["bs_diff"]
        rows = self.statement_rows("bs", 6, diff - 2, self.nfy + 1)
        ws = self.ws("bs")
        extra = (ws.cell(4, 3).value, ws.cell(5, 3).value)
        self.statement_table(rows, extra=extra)

    def periodicity_note(self):
        """Art. 5:4 §3: per-financial-year projections instead of 12/24 months, if together they cover two years."""
        if set(self.fy_months()) != {12}:
            c = self.d["company"]
            end = self.ws("cf").cell(2, self.mv.mc0 + self.mv.n - 1).value
            end = date(end.year, end.month, calendar.monthrange(end.year, end.month)[1])
            self.para(self.T("periodicity_note", start=self.f.date(c["start_date"]), end=self.f.date(end.isoformat())))

    def pl(self):
        self.h1("h_pl", legal=4)
        self.para(self.T("pl_intro"))
        self.periodicity_note()
        _, last = self.mv.rows["a9905"]
        self.statement_table(self.statement_rows("pl", 6, last, self.nfy))

    def budget(self, tmp):
        self.h1("h_budget", legal=5)
        self.para(self.T("budget_intro"))
        self.h2("h_cash_curve")
        self.doc.add_picture(str(self.chart(tmp)), width=Cm(TEXT_WIDTH))
        self.h2("h_cf_annual")
        _, last = self.mv.rows["cash"]
        ws = self.ws("cf")
        rows = []
        for r in range(4, last + 1):
            a = ws.cell(r, 1).value
            if a is None:
                continue
            vals = [ws.cell(r, 4 + j).value for j in range(self.nfy)]
            section = all(v is None for v in vals)
            bold = bool(ws.cell(r, 1).font and ws.cell(r, 1).font.b)
            if r == last:
                a = self.T("kf_cash_end")
            rows.append({"cells": [str(a)] + ([""] * self.nfy if section else self.money_cells([v or 0 for v in vals])),
                         "section": section, "total": bold and not section, "bold": bold, "indent": 0 if (section or bold) else 1})
        top, bottom = self.fy_headers()
        lw = TEXT_WIDTH - 2.6 * self.nfy
        self.table(None, rows, [lw] + [2.6] * self.nfy, header_rows=[top, bottom])
        period = self.T("vat_per_month" if int(self.param("p_VATPeriod")) == 1 else "vat_per_quarter")
        self.para(self.T("cf_note", period=period) if self.has_vat() else self.T("cf_note_novat"), size=9, italic=True, color="555555")
        self.h2("h_inc_exp")
        self.para(self.T("inc_exp_note"), size=9, italic=True, color="555555")
        self.income_expenses()

    def income_expenses(self):
        mv = self.mv
        rows, inc, exp = [], [0] * self.nfy, [0] * self.nfy

        def add(label, vals, sign, indent=1):
            if all(abs(v or 0) < 0.5 for v in vals):
                return
            rows.append({"cells": [label] + self.money_cells(vals), "indent": indent})
            for j, v in enumerate(vals):
                (inc if sign > 0 else exp)[j] += v or 0
        rows.append({"cells": [self.T("ie_income")] + [""] * self.nfy, "section": True})
        for sid, name, _, r_rev in self.lists.get("streams", []):
            add(name, self.fy_vals("revenue", r_rev), 1)
        add(self.T("ie_grants"), mv.annual("a74"), 1)
        rows.append({"cells": [self.T("ie_total_income")] + self.money_cells(inc), "total": True, "bold": True})
        rows.append({"cells": [self.T("ie_expenses")] + [""] * self.nfy, "section": True})
        add(self.T("ie_cogs"), mv.annual("a60"), -1)
        _, r_ovc = mv.rows["ovc"]
        add(self.T("ie_var"), self.fy_vals("revenue", r_ovc), -1)
        _, first, last = mv.ranges["cost_lines"]
        ws = self.ws("costs")
        for r in range(first, last + 1):
            name = ws.cell(r, 1).value
            if name:
                add(str(name), self.fy_vals("costs", r), -1)
        for r in self.lists.get("staff", []):
            name = ws.cell(r, 1).value
            if name:
                add(str(name) + self.T("ie_staff_suffix"), self.fy_vals("costs", r), -1)
        add(self.T("ie_dep"), mv.annual("a630"), -1)
        _, r_int = mv.rows["loan_int"]
        add(self.T("ie_interest"), self.fy_vals("finance", r_int), -1)
        add(self.T("ie_tax"), mv.annual("a67"), -1)
        rows.append({"cells": [self.T("ie_total_expenses")] + self.money_cells(exp), "total": True, "bold": True})
        res = [i - e for i, e in zip(inc, exp)]
        for j, (a, b) in enumerate(zip(res, mv.annual("a9904"))):
            if abs(a - b) > 1:
                raise BuildError(f"income/expenses table does not add up to the P&L result in FY{j + 1}: {a:.2f} vs {b:.2f}")
        rows.append({"cells": [self.T("ie_result")] + self.money_cells(res), "total": True, "bold": True})
        top, bottom = self.fy_headers()
        lw = TEXT_WIDTH - 2.6 * self.nfy
        self.table(None, rows, [lw] + [2.6] * self.nfy, header_rows=[top, bottom])

    def chart(self, tmp):
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib.ticker import FuncFormatter
        months = [to_date(m) for m in self.mv.months()]
        cash = [v or 0 for v in self.mv.monthly("cash")]
        fig, ax = plt.subplots(figsize=(7.2, 3.0), dpi=200)
        plt.rcParams["font.family"] = "sans-serif"
        ax.plot(months, cash, color="#" + TEAL, linewidth=2)
        ax.fill_between(months, cash, 0, where=[c >= 0 for c in cash], color="#" + TEAL, alpha=0.12, interpolate=True)
        ax.fill_between(months, cash, 0, where=[c < 0 for c in cash], color="#C00000", alpha=0.25, interpolate=True)
        ax.axhline(0, color="#808080", linewidth=0.8)
        i = min(range(len(cash)), key=lambda k: cash[k])
        ax.scatter([months[i]], [cash[i]], color="#C00000" if cash[i] < 0 else "#" + TEAL, zorder=3, s=18)
        span = (max(cash) - min(min(cash), 0)) or 1
        ax.annotate(f"{self.T('chart_lowest')}: {self.f.eur(cash[i])}", (months[i], cash[i]), textcoords="offset points",
                    xytext=(8, 10), fontsize=7, color="#333333",
                    bbox={"boxstyle": "round,pad=0.2", "fc": "white", "ec": "none", "alpha": 0.8})
        ax.set_ylim(min(min(cash), 0) - 0.08 * span, max(cash) + 0.08 * span)
        ws = self.ws("cf")
        for k in range(1, len(months)):
            if ws.cell(3, self.mv.mc0 + k).value != ws.cell(3, self.mv.mc0 + k - 1).value:
                ax.axvline(months[k], color="#BBBBBB", linewidth=0.6, linestyle="--")
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: self.f.num(v, dash=False)))
        ticks = [m for m in months if m.month in (1, 7)] or months[::6]
        ax.set_xticks(ticks)
        ax.set_xticklabels([self.f.month_short(m) for m in ticks], fontsize=7)
        ax.tick_params(axis="y", labelsize=7)
        ax.set_title(self.T("chart_title"), fontsize=8.5, loc="left", color="#333333")
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        ax.grid(axis="y", color="#EEEEEE", linewidth=0.6)
        fig.tight_layout()
        out = Path(tmp) / "cash.png"
        fig.savefig(out)
        plt.close(fig)
        return out

    def assumptions(self):
        f, mv = self.f, self.mv
        ass = self.d.get("assumptions") or {}
        self.h1("h_assumptions", legal=6)
        self.para(self.T("assumptions_intro"))
        # key figures
        self.h2("kf_title")
        kf = []
        rev = mv.annual("a70")

        def pct_row(label, row):
            return {"cells": [label] + [f.pct(v / r) if r else "-" for v, r in zip(mv.annual(row), rev)], "indent": 1}
        kf.append({"cells": [self.L("pl_70")] + self.money_cells(rev), "bold": True})
        kf.append({"cells": [self.L("pl_9900")] + self.money_cells(mv.annual("a9900"))})
        kf.append(pct_row(self.T("kf_gm_pct"), "a9900"))
        kf.append({"cells": [self.L("pl_9901")] + self.money_cells(mv.annual("a9901"))})
        kf.append(pct_row(self.T("kf_op_pct"), "a9901"))
        kf.append({"cells": [self.T("kf_net")] + self.money_cells(mv.annual("a9904")), "bold": True})
        kf.append({"cells": [self.T("kf_cash_end")] + self.money_cells(mv.annual("bs_cash"))})
        kf.append({"cells": [self.T("kf_cash_min")] + self.money_cells(mv.annual("kpi_min"))})
        kf.append({"cells": [self.L("bs_1015")] + self.money_cells(mv.annual("bs_eq"))})
        kf.append({"cells": [self.T("kf_fte")] + [f.num(v, 1, dash=False) for v in mv.fy_totals("fte")]})
        top, bottom = self.fy_headers()
        lw = TEXT_WIDTH - 2.6 * self.nfy
        self.table(None, kf, [lw] + [2.6] * self.nfy, header_rows=[top, bottom])
        self.assumption_revenue()
        self.assumption_costs()
        self.assumption_staff()
        self.assumption_investments()
        self.assumption_tax()
        self.h2("h_a_years")
        years = ass.get("years") or []
        for j in range(self.nfy):
            lab, period = mv.fy_labels[j]
            self.doc.add_heading(f"{lab} ({period})", level=3)
            if j < len(years) and re.match(r"\s*(\*\*)?\s*\{\{\s*fy\.", str(years[j])):
                self.problem(f"assumptions.years[{j}] starts with the financial-year label; the heading above it already shows it")
            self.prose(years[j] if j < len(years) else "", f"assumptions.years[{j}]")
        self.h2("h_a_profit")
        self.prose(ass.get("profitability"), "assumptions.profitability")
        if ass.get("risks") or not self.draft:
            self.h2("h_a_risks")
            self.prose(ass.get("risks"), "assumptions.risks")

    def assumption_revenue(self):
        f = self.f
        streams = self.d.get("revenue_streams", [])
        self.h2("h_a_revenue")
        prows = []
        for s in streams:
            lag = s.get("payment_lag_months", 0)
            goods = []
            if s.get("goods_cost_pct"):
                goods.append(f.pct(s["goods_cost_pct"]))
            if s.get("goods_cost_per_unit"):
                goods.append(f"{f.money(s['goods_cost_per_unit'])} {self.T('per_unit')}")
            amount_only = s.get("unit", "") == "€"
            price = f.money(s["price_excl_vat"])
            for m, v in sorted((s.get("price_schedule") or {}).items()):
                price += f" → {f.money(v)} ({f.month_short(m)})"
            prows.append({"cells": [s["name"], "-" if amount_only else s.get("unit", ""),
                                    "-" if amount_only else price, f.pct(s.get("price_indexation", 0)),
                                    f.pct(s.get("vat_rate", 0.21), 0),
                                    self.T("pay_cash") if not lag else self.months(lag), " + ".join(goods) or "-"]})
        self.table([self.T("c_stream"), self.T("c_unit"), self.T("c_price"), self.T("c_price_idx"), self.T("c_vat_rate"),
                    self.T("c_pay_term"), self.T("c_goods_pct")], prows, [4.4, 2.0, 2.0, 2.0, 1.4, 2.2, 2.6], num_from=2, size=8)
        vrows, tot = [], [0] * self.nfy
        for sid, name, r_vol, r_rev in self.lists.get("streams", []):
            unit = next((s.get("unit", "") for s in streams if s["id"] == sid), "")
            vrows.append({"cells": [name] + [""] * self.nfy, "section": True})
            if unit != "€":
                vrows.append({"cells": [f"{self.T('c_volume')} ({unit})"] + [f.num(v) for v in self.fy_vals("revenue", r_vol)], "indent": 1})
            rv = self.fy_vals("revenue", r_rev)
            vrows.append({"cells": [self.T("c_revenue")] + self.money_cells(rv), "indent": 1})
            tot = [a + b for a, b in zip(tot, rv)]
        vrows.append({"cells": [self.T("total")] + self.money_cells(tot), "total": True, "bold": True})
        top, bottom = self.fy_headers()
        lw = TEXT_WIDTH - 2.6 * self.nfy
        self.table(None, vrows, [lw] + [2.6] * self.nfy, header_rows=[top, bottom])

    def freq(self, c):
        if c.get("schedule"):
            return self.T("freq_schedule")
        n = c.get("interval_months", 1)
        return self.T(f"freq_{n}") if f"freq_{n}" in TEXT else self.T("freq_n", n=n)

    def assumption_costs(self):
        f = self.f
        self.h2("h_a_costs")
        rows = []
        for c in self.d.get("cost_lines", []):
            sch = sorted(m for m, v in (c.get("schedule") or {}).items() if v)
            amt = " / ".join(f.money(v) for v in dict.fromkeys(c["schedule"][m] for m in sch)) if sch else f.money(c["amount"])
            last = to_date(self.mv.months()[-1])
            start, end = (sch[0], sch[-1] if to_date(sch[-1]) < last.replace(day=1) else None) if sch else (c.get("start"), c.get("end"))
            rows.append({"cells": [c["name"], amt, self.freq(c), f.month_short(start) if start else "",
                                   f.month_short(end) if end else "", f.pct(c.get("indexation", 0))]})
        self.table([self.T("c_cost"), self.T("c_amount"), self.T("c_frequency"), self.T("c_from"), self.T("c_until"), self.T("c_indexation")],
                   rows, [6.6, 2.2, 2.6, 1.8, 1.8, 1.6], num_from=1, size=8)

    def assumption_staff(self):
        f = self.f
        self.h2("h_a_staff")
        staff = self.d.get("staff", [])
        if not staff:
            self.para(self.T("no_staff"))
            return
        rows = []
        for s in staff:
            if s.get("employer_cost_monthly"):
                cost = f.money(s["employer_cost_monthly"]) + "*"
            else:  # same formula as the model, at the start salary
                cost = f.eur(s.get("gross_monthly", 0) * s.get("fte", 1) * self.param("p_SalMonths") / 12 * (1 + self.param("p_SSC"))
                               + s.get("extras_monthly", 0))
            rows.append({"cells": [s["role"], s.get("contract", ""), f.month_short(s["start"]), f.month_short(s["end"]) if s.get("end") else "",
                                   f.num(s.get("fte", 1), 1, dash=False), f.money(s["gross_monthly"]) if s.get("gross_monthly") else "", cost]})
        self.table([self.T("c_role"), self.T("c_contract"), self.T("c_from"), self.T("c_until"), self.T("c_fte"), self.T("c_gross"),
                    self.T("c_employer_cost")], rows, [3.8, 2.4, 1.7, 1.7, 1.2, 3.0, 2.8], num_from=4, size=8)
        if any(x.get("employer_cost_monthly") for x in staff):
            self.para(self.T("employer_cost_note"), size=9, italic=True, color="555555", after=2)
        self.para(self.T("staff_method", months=f.num(self.param("p_SalMonths"), 2, dash=False), ssc=f.pct(self.param("p_SSC")),
                         idx=f.pct(self.param("p_WageIdx"))), size=9, italic=True, color="555555")

    def assumption_investments(self):
        f = self.f
        self.h2("h_a_invest")
        inv = self.d.get("investments", [])
        if not inv:
            self.para(self.T("no_investments"))
            return
        rows = []
        for i in inv:
            life = i.get("life_years", 3) if i.get("category", 22) != 28 else 0
            name = i["name"] + (f" ({self.T('in_kind_mark')})" if i.get("in_kind") else "")
            rows.append({"cells": [name, f.month_short(i["date"]), f.num(i["amount"]),
                                   self.T("life_years", n=life) if life else self.T("not_depreciated")]})
        rows.append({"cells": [self.T("total"), "", f.num(sum(i["amount"] for i in inv)), ""], "total": True, "bold": True})
        self.table([self.T("c_investment"), self.T("c_date"), self.T("c_amount"), self.T("c_life")], rows, [8.2, 2.4, 2.8, 3.2], num_from=1)

    def assumption_tax(self):
        f = self.f
        self.h2("h_a_tax")
        std = f.pct(self.param("p_TaxStd"), 0)
        if int(self.param("p_UseRed")) == 1:
            ws = self.ws("pl")
            _, r_el = self.mv.rows["a_eligible"]
            years = [self.mv.fy_labels[j][0] for j in range(self.nfy) if ws.cell(r_el, 3 + j).value == 1]
            red, cap = f.pct(self.param("p_TaxRed"), 0), f.eur(self.param("p_TaxRedCap"))
            if self.nfy < int(self.param("p_TaxRedFrom")):
                txt = self.T("tax_red_starter", red=red, cap=cap)
            else:
                txt = self.T("tax_red_text", red=red, cap=cap, minrem=f.eur(self.param("p_TaxRedMinRem")),
                             years=", ".join(years) or self.T("tax_red_none"), std=std) + " " + self.T("tax_red_from_note")
        else:
            txt = self.T("tax_std_text", std=std)
        txt += " " + self.T("tax_losses") + " "
        txt += self.T("tax_prepay") if int(self.param("p_Prepay")) == 1 else self.T("tax_payment", n=int(self.param("p_TaxLag")))
        self.para(txt)

    def expert(self):
        e = (self.d.get("document") or {}).get("expert")
        self.h1("h_expert", legal=7)
        if not e:
            self.para(self.T("no_expert_text"))
            return
        self.para(self.T("expert_text"))
        rows = [(self.T("c_name"), e.get("name", ""))]
        for k, lab in (("firm", "f_firm"), ("itaa", "f_itaa"), ("enterprise_number", "f_enterprise_nr"), ("address", "c_address")):
            if e.get(k):
                rows.append((self.T(lab), e[k]))
        self.kv_table(rows)

    def disclaimer(self):
        self.h1("h_disclaimer")
        checked = self.f.date(self.mv.meta["rules_checked"])
        for k in ("disc_1", "disc_2", "disc_snapshot"):
            self.para(self.T(k))
        self.para(self.T("disc_3", date=checked))
        verify = json.loads(self.mv.meta.get("rules_verify") or "[]")
        if verify:
            items = [self.L(RULE_LABEL[k])[0].lower() + self.L(RULE_LABEL[k])[1:] for k in verify if k in RULE_LABEL]
            self.para(self.T("disc_verify", items="; ".join(items)))
        self.para(self.T("disc_model"))
        self.doc.add_heading(self.T("sources_title"), level=3)
        used = self.sources_used()
        for key, src in RULES["sources"].items():
            if not src["url"] or key not in used:
                continue
            p = self.doc.add_paragraph(style="List Bullet")
            p.add_run(src["name"][self.lang] + ": ")
            add_hyperlink(p, src["url"], src["url"])
            for r in p.runs:
                r.font.size = Pt(9)

    def sources_used(self):
        """Official sources behind a value or statement of this plan."""
        d = self.d
        used = {"wvv", "wib", "vat", "nbb"}
        used |= {RULES["rules"][k]["source"] for k in json.loads(self.mv.meta.get("rules_used") or "[]") if k in RULES["rules"]}
        if d.get("staff"):
            used |= {"rsz", "index"}
        lines = d.get("cost_lines", [])
        if any(c.get("id") == "company_contribution" for c in lines) or "contribution_exempt" in d["company"]:
            used.add("rsvz_amount")
        if d["company"].get("contribution_exempt"):
            used.add("rsvz_exempt")
        if any(c.get("category") == 618 or str(c.get("id", "")).startswith("dir_social") for c in lines):
            used.add("self_employed")
        return used

    def signature(self):
        docd = self.d.get("document") or {}
        self.h1("h_signature")
        place = docd.get("place") or self.missing_inline("document.place")
        self.para(self.T("signed_at", place=place, date=self.f.date(self.doc_date())), after=24)
        names = docd.get("signatories") or [x["name"] for x in self.d.get("founders") or []]
        if not names:
            names = [self.missing_inline("document.signatories")]
        t = self.doc.add_table(rows=0, cols=2)
        t.autofit = False
        for i in range(0, len(names), 2):
            row = t.add_row()
            no_split(row)
            for j, n in enumerate(names[i:i + 2]):
                c = row.cells[j]
                c.width = Cm(8.3)
                p = c.paragraphs[0]
                p.add_run(n).bold = True
                p.paragraph_format.space_after = Pt(40)
                p2 = c.add_paragraph("_" * 34)
                p2.add_run(f"\n{self.T('signature_line')}").font.size = Pt(8)
                p2.paragraph_format.space_after = Pt(24)
        fix_widths(t, [8.3, 8.3])

    def appendix(self):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21)
        sec.left_margin = sec.right_margin = Cm(1.6)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
        self.h1_no += 1
        self.doc.add_heading(self.T("h_appendix"), level=1)
        self.toc.append((1, self.T("h_appendix")))
        self.para(self.T("appendix_intro" if self.has_vat() else "appendix_intro_novat"), size=9, italic=True, color="555555")
        ws = self.ws("cf")
        _, last = self.mv.rows["cash"]
        months = self.mv.months()
        width = 29.7 - 3.2
        for start in range(0, self.mv.n, 12):
            idx = list(range(start, min(start + 12, self.mv.n)))
            rows = []
            for r in range(4, last + 1):
                a = ws.cell(r, 1).value
                if a is None:
                    continue
                vals = [ws.cell(r, self.mv.mc0 + k).value for k in idx]
                section = all(v is None for v in vals)
                bold = bool(ws.cell(r, 1).font and ws.cell(r, 1).font.b)
                rows.append({"cells": [str(a)] + ([""] * len(idx) if section else self.money_cells([v or 0 for v in vals])),
                             "section": section, "total": bold and not section, "bold": bold})
            hdr = [self.T("in_eur")] + [self.f.month_short(months[k]) for k in idx]
            lw = 6.1
            self.table(hdr, rows, [lw] + [(width - lw) / 12] * len(idx), size=6.5)

    def build(self, out):
        self.setup()
        with tempfile.TemporaryDirectory() as tmp:
            self.cover()
            self.company()
            self.founders()
            self.activity()
            self.financing()
            self.balance()
            self.pl()
            self.budget(tmp)
            self.assumptions()
            self.expert()
            self.disclaimer()
            self.signature()
            self.appendix()
            self.fill_toc()
            if self.problems and not self.draft:
                raise BuildError("document incomplete:\n- " + "\n- ".join(self.problems))
            self.doc.save(out)
        return self.problems


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan_data")
    ap.add_argument("model")
    ap.add_argument("out")
    ap.add_argument("--draft", action="store_true")
    a = ap.parse_args()
    data = json.loads(Path(a.plan_data).read_text())
    try:
        problems = PlanDocument(data, a.model, draft=a.draft).build(a.out)
    except BuildError as e:
        print(json.dumps({"status": "error", "message": str(e)}, ensure_ascii=False, indent=2))
        sys.exit(1)
    print(json.dumps({"status": "draft" if a.draft else "success", "out": a.out, "gaps": problems}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
