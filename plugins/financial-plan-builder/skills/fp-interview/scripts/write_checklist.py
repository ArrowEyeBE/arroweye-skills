#!/usr/bin/env python3
"""Write the preparation checklist (interview question 0.2) as a Word file the founders can share.

Usage: write_checklist.py <checklist.json> <out.docx>

checklist.json, already in the plan language:
  {"title": "...", "intro": "...",
   "sections": [{"heading": "...", "columns": ["Item", "Used in"], "rows": [["...", "..."], ...]}],
   "footer": "..."}
The content comes from the "Preparation checklist" in references/question-bank.md, translated by Claude.
"""
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent / "fp-model/scripts"))
import plugin_env  # noqa: E402

plugin_env.ensure()

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

TEAL = RGBColor(0x1F, 0x4E, 0x5A)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        sys.exit(0)
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    c = json.loads(Path(sys.argv[1]).read_text())
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    h = doc.add_paragraph()
    r = h.add_run(c["title"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), TEAL
    if c.get("intro"):
        doc.add_paragraph(c["intro"])
    for s in c["sections"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(s["heading"])
        r.bold, r.font.size, r.font.color.rgb = True, Pt(12), TEAL
        cols = s.get("columns") or []
        n = max(len(cols), max((len(x) for x in s["rows"]), default=1))
        t = doc.add_table(rows=0, cols=n + 1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [Cm(0.8), Cm(12.2)] + [Cm(4.0)] * (n - 1)
        if cols:
            cells = t.add_row().cells
            for i, v in enumerate([""] + cols):
                cells[i].text = v
                shade(cells[i], "1F4E5A")
                for run in cells[i].paragraphs[0].runs:
                    run.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        for row in s["rows"]:
            cells = t.add_row().cells
            cells[0].text = "☐"
            for i, v in enumerate(row):
                cells[i + 1].text = str(v)
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[min(i, len(widths) - 1)]
    if c.get("footer"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        r = p.add_run(c["footer"])
        r.italic, r.font.size = True, Pt(9)
    doc.save(sys.argv[2])
    print(json.dumps({"status": "success", "out": sys.argv[2]}))


if __name__ == "__main__":
    main()
