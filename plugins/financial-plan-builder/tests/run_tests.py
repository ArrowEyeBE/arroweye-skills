#!/usr/bin/env python3
"""Test suite for the financial-plan model.

For every case in tests/cases/*/plan-data.json (and the private regression cases in
reference/tests/*/, when present): build the model in NL, FR and EN, recalculate it,
and assert zero formula errors, a balanced balance sheet and the expected check results.
A case with `expected` values is compared within its own `tolerance` (default 0.01 EUR);
a case with `expect_review` must produce those reviewer findings.

Usage: python3 tests/run_tests.py [--lang nl]
"""
import argparse
import copy
import json
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "skills/fp-model/scripts"
DOC_SCRIPTS = ROOT / "skills/fp-document/scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(DOC_SCRIPTS))
from check_model import legal_rules, summarise  # noqa: E402
from validate_plan_data import validate  # noqa: E402
from model_io import ModelValues  # noqa: E402
from build_document import Fmt  # noqa: E402



def run(cmd):
    return subprocess.run([sys.executable, *map(str, cmd)], capture_output=True, text=True)


def build(data, tmp, tag):
    pd = tmp / f"{tag}.json"
    pd.write_text(json.dumps(data, ensure_ascii=False))
    raw, val = tmp / f"{tag}.xlsx", tmp / f"{tag}.values.xlsx"
    r = run([SCRIPTS / "build_model.py", pd, raw])
    if r.returncode:
        return None, f"build failed: {r.stderr[-800:]}"
    r = run([SCRIPTS / "recalc.py", raw, val])
    res = json.loads(r.stdout or "{}")
    if res.get("status") != "success":
        return None, f"formula errors: {json.dumps(res.get('errors'))[:800]}"
    return val, None


def build_doc(data, model, tmp, tag, draft=False):
    pd = tmp / f"{tag}.doc.json"
    pd.write_text(json.dumps(data, ensure_ascii=False))
    out = tmp / f"{tag}.docx"
    r = run([DOC_SCRIPTS / "build_document.py", pd, model, out] + (["--draft"] if draft else []))
    try:
        res = json.loads(r.stdout)
    except json.JSONDecodeError:
        res = {"status": "crash", "message": r.stderr[-800:]}
    return r.returncode, res, out


def months_in_first_fy(c):
    s, e = date.fromisoformat(c["start_date"]), date.fromisoformat(c["first_fy_end"])
    return (e.year - s.year) * 12 + e.month - s.month + 1


def doc_text(path):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def check_document(name, data, model, tmp):
    """Final document must build, contain every legal chapter and show the model's numbers."""
    failures = []
    code, res, out = build_doc(data, model, tmp, name)
    if code:
        return [f"document: {res.get('message', res)}"]
    text = doc_text(out)
    lang = data["meta"]["language"]
    from doc_i18n import TEXT
    for key in ("h_company", "h_founders", "h_activity", "h_financing", "h_balance", "h_pl", "h_budget", "h_assumptions",
                "h_expert", "h_disclaimer", "h_signature", "h_appendix"):
        if TEXT[key][lang] not in text:
            failures.append(f"document: chapter missing: {TEXT[key][lang]}")
    # art. 5:4 §2: every one of the seven required elements is referenced
    for n in range(1, 8):
        if TEXT["legal_ref"][lang].format(n=n) not in text:
            failures.append(f"document: art. 5:4 §2 {n}° not referenced")
    # §3: projections per financial year are only allowed with a note, when the first year is not 12 months
    note = TEXT["periodicity_note"][lang].split("{")[0]
    first_fy = months_in_first_fy(data["company"])
    if (note in text) != (first_fy != 12):
        failures.append(f"document: periodicity note {'missing' if first_fy != 12 else 'unexpected'} (first FY {first_fy} months)")
    has_kind = any(i.get("in_kind") for i in data.get("investments", []))
    if (TEXT["inkind_note"][lang][:40] in text) != has_kind:
        failures.append("document: art. 5:7 note on contributions in kind wrong")
    if not (data.get("document") or {}).get("expert") and TEXT["no_expert_text"][lang] not in text:
        failures.append("document: no statement that no external expert assisted")
    for marker in ("[?", TEXT["missing"][lang].split("{")[0]):
        if marker in text:
            failures.append(f"document: contains '{marker}'")
    mv, f = ModelValues(model), Fmt(lang)
    for row in ("a70", "a9904", "bs_cash", "bs_eq"):
        for v in mv.annual(row):
            if f.num(v) not in text:
                failures.append(f"document: {row} value {f.num(v)} not found")
    return failures


def check_case(name, data, tmp, expect_status=None):
    failures = [f"invalid plan-data: {e['path']}: {e['message']}" for e in validate(data)["errors"]]
    val, err = build(data, tmp, name)
    if err:
        return [err]
    s = summarise(val)
    if any(abs(d or 0) > 0.01 for d in s["balance_sheet_difference"]):
        failures.append(f"balance sheet does not balance: {s['balance_sheet_difference']}")
    status = {c["check"]: c["status"] for c in s["checks"]}
    for ck, st in (expect_status or {}).items():
        if status.get(ck) != st:
            failures.append(f"{ck}: expected {st}, got {status.get(ck)}")
    if not expect_status or "ck_cash" not in expect_status:
        if status.get("ck_cash") == "error":
            failures.append(f"cash goes negative: {s['lowest_cash']}")
    if data.get("activity") and not any(st == "error" for st in status.values()):
        failures += check_document(name, data, val, tmp)
    exp = data.get("expected")
    if exp:
        mv = ModelValues(val)
        for group in ("pl", "bs"):
            for row, values in exp.get(group, {}).items():
                got = mv.annual(row)
                tol = (data.get("tolerance") or {}).get(row, 0.01)
                for i, (g, e) in enumerate(zip(got, values)):
                    if abs(g - e) > tol:
                        failures.append(f"{row} FY{i + 1}: {g:,.2f} vs expected {e:,.2f} (tolerance {tol})")
    return failures


def check_rules_override(data, tmp):
    val, err = build(data, tmp, "rules-override")
    if err:
        return [err]
    failures = []
    lr = summarise(val)["legal_rules"]
    if "reduced_rate_min_director_pay" in lr["to_verify"]:
        failures.append("overridden rule still listed as to verify")
    mv = ModelValues(val)
    ws = mv.wb[mv.sheets["params"]]
    row = next((r for r in ws.iter_rows(values_only=True) if r[1] == 48000), None)
    if not row or row[3] != "test source" or row[4] != "10/01/2027":
        failures.append(f"override source/date not shown in settings: {row}")
    return failures


def check_document_guards(base, tmp):
    """The document refuses to go final with a cash shortfall, missing text or an unknown placeholder."""
    failures = []
    d = copy.deepcopy(base)
    d["financing"]["loans"][0]["amount"] = 90000
    val, err = build(d, tmp, "guard-cash")
    if err:
        return [err]
    code, res, _ = build_doc(d, val, tmp, "guard-cash")
    if code == 0 or "ck_cash" not in res.get("message", ""):
        failures.append(f"final document built despite negative cash: {res}")
    code, res, _ = build_doc(d, val, tmp, "guard-cash-draft", draft=True)
    if code:
        failures.append(f"draft document with negative cash failed: {res}")
    d = copy.deepcopy(base)
    d["assumptions"]["profitability"] = ""
    d["assumptions"]["years"][0] += " {{revenue.9}} {{no_such_figure.1}}"
    val, err = build(d, tmp, "guard-text")
    if err:
        return failures + [err]
    code, res, _ = build_doc(d, val, tmp, "guard-text")
    msg = res.get("message", "")
    for needle in ("assumptions.profitability", "revenue.9", "no_such_figure.1"):
        if code == 0 or needle not in msg:
            failures.append(f"final document did not report '{needle}': {res}")
    code, res, _ = build_doc(d, val, tmp, "guard-text-draft", draft=True)
    if code or len(res.get("gaps", [])) != 3:
        failures.append(f"draft should build and list 3 gaps: {res}")
    # "there are no grants" may only be printed when the owner confirmed it
    d = copy.deepcopy(base)
    d["financing"]["grants"] = []
    d["financing"]["equity"][0]["amount"] += 10000  # keep cash positive, so only the confirmation is missing
    val, err = build(d, tmp, "guard-none")
    if err:
        return failures + [err]
    code, res, _ = build_doc(d, val, tmp, "guard-none")
    if code == 0 or "confirmed_none" not in res.get("message", ""):
        failures.append(f"final document states 'no grants' without the owner's confirmation: {res}")
    return failures


def check_horizon(base, tmp):
    """Art. 5:4: the projections must reach at least two years after incorporation."""
    failures = []
    for start, nfy, expected in (("2027-01-04", 2, "error"), ("2027-01-01", 2, "ok"), ("2027-01-04", 3, "ok")):
        d = copy.deepcopy(base)
        d.pop("expected", None)
        d["company"].update(start_date=start, first_fy_end="2027-12-31", fiscal_years=nfy)
        val, err = build(d, tmp, f"horizon-{start}-{nfy}")
        if err:
            failures.append(err)
            continue
        got = {c["check"]: c["status"] for c in summarise(val)["checks"]}.get("ck_horizon")
        if got != expected:
            failures.append(f"start {start}, {nfy} financial years: ck_horizon {got}, expected {expected}")
    return failures


def review(data, model, tmp, tag, docx=None):
    pd = tmp / f"{tag}.review.json"
    pd.write_text(json.dumps(data, ensure_ascii=False))
    r = run([DOC_SCRIPTS / "review_plan.py", pd, model] + ([docx] if docx else []))
    try:
        return json.loads(r.stdout)["findings"]
    except (json.JSONDecodeError, KeyError):
        return [{"severity": "crash", "code": "crash", "where": "", "message": r.stderr[-800:]}]


def check_review(base, tmp, private=()):
    """The reviewer's mechanical part: a clean plan passes, planted contradictions are caught."""
    failures = []
    d = copy.deepcopy(base)
    d["meta"]["language"] = "nl"
    val, err = build(d, tmp, "review")
    if err:
        return [err]
    code, res, docx = build_doc(d, val, tmp, "review")
    if code:
        return [f"document: {res}"]
    allowed = {"model_ck_buffer"}
    bad = [f for f in review(d, val, tmp, "clean", docx) if f["severity"] in ("error", "warning", "crash") and f["code"] not in allowed]
    if bad:
        failures.append(f"clean plan has findings: {bad}")
    # every other public case, in its own language, must review clean too
    for p in sorted((ROOT / "tests/cases").glob("*/plan-data.json")):
        c = json.loads(p.read_text())
        if p.parent.name == "bakkerij-de-korst" or not c.get("activity"):
            continue
        v, err = build(c, tmp, f"review-{p.parent.name}")
        if err:
            failures.append(err)
            continue
        code, res, dx = build_doc(c, v, tmp, f"review-{p.parent.name}")
        bad = [f for f in review(c, v, tmp, p.parent.name, None if code else dx)
               if f["severity"] in ("error", "warning", "crash") and f["code"] not in allowed]
        if code or bad:
            failures.append(f"{p.parent.name}: {res if code else bad}")
    planted = [
        ("typed amount", lambda x: x["assumptions"]["years"].__setitem__(0, x["assumptions"]["years"][0] + " De winst bedraagt 25.000 euro."),
         "untraced_number", "assumptions.years[0]"),
        ("hire not in model", lambda x: x["assumptions"]["years"].__setitem__(2, x["assumptions"]["years"][2] + " Er wordt een tweede bakker aangeworven."),
         "hire_not_in_model", "assumptions.years[2]"),
        ("denies loan", lambda x: x["document"].__setitem__("financing_note", "De bakkerij wordt volledig met eigen middelen gefinancierd, zonder lening."),
         "denies_existing_loan", "document.financing_note"),
        ("wrong language", lambda x: x["assumptions"].__setitem__("risks", "- Fewer customers than expected: the founders will reduce the opening hours and the student will work fewer weekends."),
         "wrong_language", "assumptions.risks"),
    ]
    for label, change, want, where in planted:
        x = copy.deepcopy(d)
        change(x)
        got = review(x, val, tmp, "planted")
        if not any(f["code"] == want and f["where"] == where for f in got):
            failures.append(f"planted '{label}' not caught: {[(f['code'], f['where']) for f in got]}")
    x = copy.deepcopy(d)
    x["financing"]["loans"] = []
    x["assumptions"]["years"][1] += " Een bijkomende lening van de bank financiert een tweede oven."
    val2, err = build(x, tmp, "review-noloan")
    if err:
        return failures + [err]
    if not any(f["code"] == "loan_not_in_model" for f in review(x, val2, tmp, "noloan")):
        failures.append("planted 'loan not in model' not caught")
    for i, p in enumerate(private):
        d = json.loads(p.read_text())
        if not d.get("expect_review"):
            continue
        val, err = build(d, tmp, f"review-private-{i}")
        if err:
            return failures + [err]
        got = review(d, val, tmp, f"private-{i}")
        for want in d["expect_review"]:
            if not any(f["code"] == want["code"] and f["where"] == want["where"] for f in got):
                failures.append(f"private case {i + 1}: expected finding {want['code']} at {want['where']} not found")
    return failures


def check_validator(base):
    """Planted input errors are reported on the right path; a valid case has no errors."""
    def drop_loans(x):
        x["financing"]["loans"] = []
    def short_horizon(x):
        x["company"]["fiscal_years"] = 2
    def volume_before_start(x):
        x["revenue_streams"][0]["volumes"]["2026-12"] = 100
    def grace_too_long(x):
        x["financing"]["loans"][0]["grace_months"] = 84
    def inkind_vat(x):
        next(i for i in x["investments"] if i.get("in_kind"))["vat_rate"] = 0.21
    def euro_price(x):
        x["revenue_streams"][0]["unit"] = "€"
    def no_language(x):
        x["meta"]["language"] = "de"
    def bad_vat(x):
        x["cost_lines"][0]["vat_rate"] = 0.19
    def wrong_confirmation(x):
        x["financing"]["confirmed_none"] = ["grants"]
    cases = [(drop_loans, "financing.confirmed_none"), (short_horizon, "company.fiscal_years"),
             (volume_before_start, "revenue_streams[0].volumes['2026-12']"), (grace_too_long, "financing.loans[0].grace_months"),
             (inkind_vat, "investments[5].vat_rate"), (euro_price, "revenue_streams[0].price_excl_vat"),
             (no_language, "meta.language"), (bad_vat, "cost_lines[0].vat_rate"), (wrong_confirmation, "financing.confirmed_none")]
    failures = [f"valid case has errors: {validate(base)['errors']}"] if validate(base)["errors"] else []
    for change, path in cases:
        x = copy.deepcopy(base)
        change(x)
        errs = validate(x)["errors"]
        if not any(e["path"] == path for e in errs):
            failures.append(f"{change.__name__}: no error on {path}: {errs}")
    return failures


def check_checklist(tmp):
    spec = {"title": "Checklist", "intro": "Test", "footer": "f",
            "sections": [{"heading": "Documenten", "columns": ["Item", "Gebruikt in"], "rows": [["Ontwerp statuten", "1.4, 1.5"]]}]}
    (tmp / "cl.json").write_text(json.dumps(spec))
    r = run([ROOT / "skills/fp-interview/scripts/write_checklist.py", tmp / "cl.json", tmp / "cl.docx"])
    if r.returncode or "Ontwerp statuten" not in doc_text(tmp / "cl.docx"):
        return [f"checklist not written: {r.stderr[-400:]}"]
    return []


def check_env_switch(tmp):
    """check_env.py offers the packages, and the scripts switch to the plugin's own environment."""
    import os
    import site
    env_script = ROOT / "skills/belgian-financial-plan/scripts/check_env.py"
    bin_py = "Scripts/python.exe" if os.name == "nt" else "bin/python"
    failures = []
    # a Python without the packages, and a plugin environment that has them (via a .pth to this Python's)
    for name in ("bare", "full"):
        subprocess.run([sys.executable, "-m", "venv", "--without-pip", tmp / name], check=True)
    full_site = next((tmp / "full").glob("lib*/**/site-packages"), None) or (tmp / "full/Lib/site-packages")
    dirs = [*site.getsitepackages(), site.getusersitepackages()]
    (full_site / "parent.pth").write_text("\n".join(dirs) + "\n")
    bare = str(tmp / "bare" / bin_py)

    def with_venv(venv):
        return {**os.environ, "FINANCIAL_PLAN_VENV": str(venv)}

    def env_json(venv):
        r = subprocess.run([bare, env_script], capture_output=True, text=True, env=with_venv(venv))
        return json.loads(r.stdout)

    out = env_json(tmp / "none")
    pkg = [m for m in out["missing"] if m["what"].startswith("Python packages")]
    if out["status"] != "missing" or not pkg or not pkg[0]["claude_may_run"] or "pip install" not in pkg[0]["install"]:
        failures.append(f"missing packages not offered: {out}")
    out = env_json(tmp / "full")
    if out["python"] != str(tmp / "full" / bin_py) or any(m["what"].startswith("Python packages") for m in out["missing"]):
        failures.append(f"plugin environment not recognised: {out}")
    data = json.loads((ROOT / "tests/cases/bakkerij-de-korst/plan-data.json").read_text())
    (tmp / "env.json").write_text(json.dumps(data))
    r = subprocess.run([bare, SCRIPTS / "build_model.py", tmp / "env.json", tmp / "env.xlsx"], capture_output=True,
                       text=True, env=with_venv(tmp / "full"))
    if r.returncode or not (tmp / "env.xlsx").exists():
        failures.append(f"build_model did not switch to the plugin environment: {r.stderr[-300:]}")
    return failures


def check_privacy(private):
    """No name or address from a private regression case may appear in the public files."""
    words = set()
    for p in private:
        d = json.loads(p.read_text())
        words |= {d["company"]["name"], d["company"]["address"], p.parent.name}
        words |= {f.get(k) for f in d.get("founders", []) for k in ("name", "address")}
        words |= {(d.get("document") or {}).get("expert", {}).get(k) for k in ("name", "itaa", "enterprise_number")}
        words |= set((d.get("privacy_terms") or []))
    words = {w for w in words if w and len(w) > 4}
    # the marketplace file sits in this plugin (standalone) or two levels up (arroweye-plugins monorepo)
    market = next((c for c in (ROOT / ".claude-plugin/marketplace.json", ROOT.parent.parent / ".claude-plugin/marketplace.json") if c.exists()), None)
    allowed = {ROOT / ".claude-plugin/plugin.json", ROOT / ".claude-plugin/marketplace.json"}  # the plugin's author
    # the chosen author and marketplace name may appear elsewhere (README install line, licence); nothing else may
    own = [json.loads((ROOT / ".claude-plugin/plugin.json").read_text())["author"]["name"]]
    if market:
        own.append(json.loads(market.read_text())["name"])
    hits = []
    for f in ROOT.rglob("*"):
        rel = f.relative_to(ROOT)
        if not f.is_file() or rel.parts[0] in ("reference", ".git") or "out" in rel.parts or f in allowed:
            continue
        if f.suffix not in (".md", ".py", ".json", ".txt", ".yaml", ".yml"):
            continue
        text = f.read_text(errors="ignore").lower()
        for o in own:
            text = text.replace(o.lower(), " ").replace(o.split(" (")[0].lower(), " ")
        hits += [f"{rel}: '{w[:12]}…'" for w in words if w.lower() in text]
    return hits


def check_management(base, tmp):
    """Report of the first real run (management company): every fix stays fixed."""
    failures = []
    sys.path.insert(0, str(ROOT / "skills/fp-interview/scripts"))
    from cost_defaults import propose
    from doc_i18n import TEXT
    if not any(w["path"] == "company.name" for w in validate(base)["warnings"]):
        failures.append("validator: legal form in company.name not flagged")
    val, err = build(base, tmp, "mgmt")
    if err:
        return [err]
    mv = ModelValues(val)
    if [round(v) for v in mv.annual("a70")] != [55000, 54000, 54000]:
        failures.append(f"price_schedule: revenue {mv.annual('a70')}")
    tax, pbt = mv.annual("a67"), mv.annual("a9903")
    if abs(tax[0] - 0.2 * pbt[0]) > 0.01:
        failures.append(f"reduced rate not applied in the first taxable period: tax {tax[0]} on {pbt[0]}")
    if json.loads(mv.meta["rules_verify"]):
        failures.append(f"rules to verify listed although unused: {mv.meta['rules_verify']}")
    eq, res = mv.annual("bs_eq"), mv.annual("a9904")
    if abs(eq[1] - (eq[0] + res[1] - 6000)) > 0.01 or abs(eq[0] - (3000 + res[0] - 4000)) > 0.01:
        failures.append(f"dividends not deducted from equity: {eq}")
    code, res_doc, out = build_doc(base, val, tmp, "mgmt")
    if code:
        return failures + [f"document: {res_doc}"]
    text = doc_text(out)
    for want in ("Claes Management BV", TEXT["purpose_status"]["nl"][:22], TEXT["tax_red_starter"]["nl"].split(". ")[1][:40],
                 TEXT["div_intro"]["nl"][:40], TEXT["disc_snapshot"]["nl"][:40], "€ 4.000 → € 4.500"):
        if want not in text:
            failures.append(f"document: '{want}' missing")
    for bad in ("BV BV", TEXT["disc_verify"]["nl"][:30]):
        if bad in text:
            failures.append(f"document: contains '{bad}'")
    # placeholders: a unit next to one stops the final build; :abs and :num work
    for change, want in ((lambda x: x["assumptions"].__setitem__("profitability", x["assumptions"]["profitability"] + " Samen {{revenue.2}} euro."), "already includes"),
                         (lambda x: x["assumptions"].__setitem__("profitability", x["assumptions"]["profitability"] + " Marge {{gm_pct.2}}%."), "already includes"),
                         (lambda x: x["assumptions"]["years"].__setitem__(0, "**{{fy.1}}.** " + x["assumptions"]["years"][0]), "financial-year label")):
        x = copy.deepcopy(base)
        change(x)
        c, r, _ = build_doc(x, val, tmp, "mgmt-ph")
        if not c or want not in r.get("message", ""):
            failures.append(f"placeholder guard '{want}' did not stop the build: {r}")
    x = copy.deepcopy(base)
    x["assumptions"]["profitability"] += " Verlies {{net_result.1:abs}}, omzet {{revenue.2:num}} en marge {{gm_pct.2:num}} procent."
    c, r, out2 = build_doc(x, val, tmp, "mgmt-mod")
    t2 = doc_text(out2) if not c else ""
    if c or "omzet 54.000 en marge 32,6 procent" not in t2:
        failures.append(f"placeholder modifiers: {r if c else 'text not found'}")
    # a dividend larger than the equity breaks art. 5:142 and stops the final document
    x = copy.deepcopy(base)
    x["financing"]["dividends"][0]["amount"] = 40000
    v2, err = build(x, tmp, "mgmt-div")
    if err:
        return failures + [err]
    if {c["check"]: c["status"] for c in summarise(v2)["checks"]}.get("ck_dividend") != "error":
        failures.append("ck_dividend: excessive dividend not flagged")
    # cost defaults for a management company
    props = {p["id"]: p for p in propose(base)["proposals"]}
    bare = copy.deepcopy(base)
    bare["cost_lines"] = [c for c in bare["cost_lines"] if c["id"] not in ("dir_social", "company_contribution")]
    bp = {p["id"]: p for p in propose(bare)["proposals"]}
    for k in ("marketing", "it_software", "liability_insurance", "professional_indemnity"):
        if k in props:
            failures.append(f"cost defaults: '{k}' proposed for a management company")
    if "dir_social" not in bp or not 1000 < bp["dir_social"]["amount"] < 1300:
        failures.append(f"cost defaults: dir_social {bp.get('dir_social', {}).get('amount')}")
    if bp.get("company_contribution", {}).get("entry", {}).get("start") != "2029-01":
        failures.append(f"cost defaults: company contribution should start after the exemption, got {bp.get('company_contribution')}")
    return failures


def check_staleness():
    meta = {"rules_checked": "2026-10-15", "rules_review_days": 180, "rules_verify": "[]"}
    cases = [(date(2026, 11, 1), False), (date(2027, 1, 2), True), (date(2026, 12, 31), False)]
    meta2 = {**meta, "rules_checked": "2026-01-05"}
    cases2 = [(date(2026, 7, 1), False), (date(2026, 7, 10), True)]
    failures = []
    for m, cs in ((meta, cases), (meta2, cases2)):
        for today, stale in cs:
            if legal_rules(m, today)["stale"] != stale:
                failures.append(f"checked {m['rules_checked']}, today {today}: expected stale={stale}")
    return failures


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--lang", choices=["nl", "fr", "en"], help="only test this language")
    args = ap.parse_args()
    langs = [args.lang] if args.lang else ["nl", "fr", "en"]
    cases = sorted((ROOT / "tests/cases").glob("*/plan-data.json")) + sorted((ROOT / "reference/tests").glob("*/plan-data.json"))
    results = []
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        for p in cases:
            base = json.loads(p.read_text())
            for lang in langs:
                d = copy.deepcopy(base)
                d["meta"]["language"] = lang
                name = f"{p.parent.name}-{lang}"
                results.append((name, check_case(name, d, tmp, base.get("expect_status"))))
        # negative test: too little financing must be flagged, never fixed silently
        bakery = ROOT / "tests/cases/bakkerij-de-korst/plan-data.json"
        if bakery.exists():
            d = json.loads(bakery.read_text())
            d["financing"]["loans"][0]["amount"] = 90000
            d.pop("expected", None)
            results.append(("bakery-shortfall", check_case("bakery-shortfall", d, tmp, {"ck_cash": "error"})))
        # legal rules: an overridden rule leaves the 'to verify' list and shows its own source
        if bakery.exists():
            d = json.loads(bakery.read_text())
            d["rules"] = {**d.get("rules", {}), "reduced_rate_min_director_pay": 48000}
            d["rules_sources"] = {"reduced_rate_min_director_pay": {"source": "test source", "checked": "2027-01-10"}}
            results.append(("rules-override", check_rules_override(d, tmp)))
    results.append(("rules-staleness", check_staleness()))
    if bakery.exists():
        with tempfile.TemporaryDirectory() as t:
            results.append(("document-guards", check_document_guards(json.loads(bakery.read_text()), Path(t))))
            results.append(("horizon-2-years", check_horizon(json.loads(bakery.read_text()), Path(t))))
            results.append(("validator", check_validator(json.loads(bakery.read_text()))))
            results.append(("checklist-docx", check_checklist(Path(t))))
            results.append(("environment-switch", check_env_switch(Path(t))))
            mgmt = ROOT / "tests/cases/beheer-claes-management/plan-data.json"
            results.append(("management-company", check_management(json.loads(mgmt.read_text()), Path(t))))
            private = sorted((ROOT / "reference/tests").glob("*/plan-data.json"))
            results.append(("reviewer-signals", check_review(json.loads(bakery.read_text()), Path(t), private)))
            if private:
                results.append(("privacy", check_privacy(private)))
    bad = 0
    for name, fails in results:
        print(f"{'PASS' if not fails else 'FAIL'}  {name}")
        for f in fails:
            print(f"      - {f}")
        bad += bool(fails)
    print(f"\n{len(results) - bad}/{len(results)} passed")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
